"""Script de Carga da Tabela de Bairros x Áreas Fiscais.

Lê a planilha oficial 'Bairros de São José do Rio Preto.xlsx', aplica o algoritmo
de normalização de nomes (NFKD, sem acentos, abreviações expandidas) e popula a
tabela `bairros_areas` no banco de dados central `triagem_vre.db`.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.core.config import BASE_DIR, CONFIGURACOES_DIR, TRIAGEM_DB_PATH
from src.core.database import db_transaction, get_triagem_db, initialize_schema_if_needed
from src.core.normalizer import normalizar_bairro


def localizar_planilha_bairros(caminho_especifico: Path | str | None = None) -> Path:
    """Localiza a planilha de bairros da cidade em diretórios padrão."""
    if caminho_especifico:
        p = Path(caminho_especifico)
        if p.exists():
            return p

    candidatos = [
        BASE_DIR / "database" / "dados_referencia" / "Bairros de São José do Rio Preto.xlsx",
        CONFIGURACOES_DIR / "database" / "dados_referencia" / "Bairros de São José do Rio Preto.xlsx",
        BASE_DIR / "database" / "Bairros de São José do Rio Preto.xlsx",
        PROJECT_ROOT / "database" / "dados_referencia" / "Bairros de São José do Rio Preto.xlsx",
        PROJECT_ROOT / "Bairros de São José do Rio Preto.xlsx",
    ]

    for cand in candidatos:
        if cand.exists():
            return cand

    for match in BASE_DIR.rglob("*.xlsx"):
        if "BAIRRO" in match.name.upper():
            return match

    raise FileNotFoundError("Planilha de Bairros de São José do Rio Preto não encontrada.")


def carregar_bairros_areas(
    caminho_xlsx: Path | str | None = None,
    db_path: Path | str | None = None,
) -> int:
    """Lê a planilha de bairros, normaliza e persiste em bairros_areas."""
    arquivo_xlsx = localizar_planilha_bairros(caminho_xlsx)
    conn = get_triagem_db(db_path)
    initialize_schema_if_needed(conn, is_cep=False)

    print(f"[*] Abrindo planilha de bairros: {arquivo_xlsx.name}")
    wb = openpyxl.load_workbook(str(arquivo_xlsx), data_only=True)
    ws = wb.active

    bairros_unicos: dict[str, str] = {}

    for row_idx in range(2, ws.max_row + 1):
        nome_bairro = ws.cell(row=row_idx, column=1).value
        classif_area = ws.cell(row=row_idx, column=2).value

        if not nome_bairro or not classif_area:
            continue

        bairro_norm = normalizar_bairro(str(nome_bairro))
        area_limpa = str(classif_area).strip().upper()

        if bairro_norm and area_limpa:
            # Em caso de duplicidade de chave normalizada, mantém a primeira ocorrência
            if bairro_norm not in bairros_unicos:
                bairros_unicos[bairro_norm] = area_limpa

    registros = [(b, a) for b, a in bairros_unicos.items()]

    with db_transaction(conn) as cur:
        cur.execute("DELETE FROM bairros_areas;")
        cur.executemany(
            """
            INSERT INTO bairros_areas (bairro_normalizado, area)
            VALUES (?, ?)
            ON CONFLICT(bairro_normalizado) DO UPDATE SET
                area = excluded.area;
            """,
            registros,
        )

    print("[✓] Carga de Bairros x Áreas concluída com sucesso!")
    print(f"    - Bairros normalizados únicos inseridos: {len(registros)}")

    conn.close()
    return len(registros)


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else None
    carregar_bairros_areas(caminho)
