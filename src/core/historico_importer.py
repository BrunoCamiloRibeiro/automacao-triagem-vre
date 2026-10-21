"""Módulo de Importação de Histórico de Fiscais, CEPs e Backlog a partir de Planilhas Excel Antigas.

Especialmente projetado para processar planilhas reais e legadas da fiscalização,
que costumam conter irregularidades estruturais:
- Linhas duplicadas (mesmo CNPJ e fiscal repetidos dezenas de vezes).
- Cabeçalhos secundários repetidos no meio da planilha decorrentes de cópias manuais.
- Formatos heterogêneos de data (BR, ISO, datetime nativo, número serial do Excel, timestamps).
- CNPJs em formato float (.0), com espaços, pontuações ou sem máscara.
- Textos nulos ou indicativos de pendência no lugar do nome do fiscal (ex: 'N/A', 'Pendente', '-').

Processamento e Deduplicação Inteligente:
1. Deduplicação em memória antes do envio ao banco SQLite:
   - Para pares (CNPJ x Fiscal): retém a data de atendimento mais recente.
   - Para pares (CEP x Área): retém a associação e enriquece com bairro de referência.
   - Para protocolos: consolida dados cadastrais no `vre_backlog` com tratamento canônico de status.
2. Persistência atômica via `INSERT ... ON CONFLICT DO UPDATE` (UPSERT).
3. Registro de competências mensais descobertas em `controle_competencias`.
4. Reconciliação imediata de áreas e fiscais no backlog ativo.
"""

from __future__ import annotations

import datetime
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl

from src.core.cep_service import reconciliar_areas_backlog, registrar_area_cep
from src.core.config import (
    BASE_DIR,
    PLANILHAS_ANTERIORES_DIR,
)
from src.core.database import (
    db_transaction,
    get_cep_db,
    get_triagem_db,
    initialize_schema_if_needed,
)
from src.core.normalizer import (
    formatar_cnpj,
    normalizar_bairro,
    normalizar_cep,
    normalizar_cnpj,
)

# Conjunto estático de termos inválidos que não representam fiscais reais
TERMOS_INVALIDOS_FISCAL: set[str] = {
    "",
    "-",
    "--",
    "---",
    ",",
    ";",
    "NONE",
    "NULO",
    "NULL",
    "N/A",
    "NA",
    "N / A",
    "#N/A",
    "#REF!",
    "A DEFINIR",
    "PENDENTE",
    "DISTRIBUIR",
    "DISTRIBUIDO",
    "DISTRIBUÍDO",
    "SEM FISCAL",
    "SEM ATRIBUICAO",
    "DEFERIDO",
    "INDEFERIDO",
    "CANCELADO",
    "CASSADO",
    "FISCAL",
    "CNPJ",
    "STATUS",
    "PROTOCOLO",
}

# Dicionário de canonicidade de status para exibição profissional
STATUS_CANONICOS: dict[str, str] = {
    "DEFERIDO": "Deferido",
    "INDEFERIDO": "Indeferido",
    "CANCELADO": "Cancelado",
    "CASSADO": "Cassado",
    "EM ANDAMENTO": "Em Andamento",
}


def normalizar_data_atendimento(val: Any, data_fallback: str) -> str:
    """Harmoniza datas nos mais diversos formatos em uma string ISO 'AAAA-MM-DD'.

    Suporta:
    - Objetos `datetime.date` e `datetime.datetime`.
    - Números seriais do Excel (intervalo 30.000 a 65.000, calculados a partir de 1899-12-30).
    - Strings brasileiras 'DD/MM/AAAA' ou 'DD/MM/AA'.
    - Strings ISO 'AAAA-MM-DD'.
    - Timestamps compostos com horário (ex: '15/06/2024 10:30:00').
    - Valores nulos ou inválidos retornam a data informada em `data_fallback`.

    Args:
        val: Valor bruto contendo a data a ser normalizada.
        data_fallback: Data padrão em formato ISO caso o valor seja nulo/inválido.

    Returns:
        str: Data harmonizada no formato ISO 'AAAA-MM-DD'.

    Complexity:
        Tempo: O(1) - parsing de string e operações aritméticas simples.
        Espaço: O(1).
    """
    if val is None or val == "":
        return data_fallback

    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")

    # Número serial do Excel (ex: 30000 a 65000)
    if isinstance(val, (int, float)):
        try:
            if 30000 <= val <= 65000:
                base_date = datetime.date(1899, 12, 30)
                return (base_date + datetime.timedelta(days=int(val))).strftime("%Y-%m-%d")
        except Exception:
            pass

    s = str(val).strip()
    if not s or s.upper() in ("NONE", "NULL", "#N/A", "#REF!"):
        return data_fallback

    # Se contiver timestamp ou hora (ex: '15/06/2024 10:30:00'), isola a porção da data
    if " " in s:
        s = s.split(" ")[0].strip()

    # Formato brasileiro DD/MM/AAAA ou DD/MM/AA
    if "/" in s:
        partes = s.split("/")
        if len(partes) == 3:
            d, m, y = partes[0].zfill(2), partes[1].zfill(2), partes[2]
            if len(y) == 2:
                y = "20" + y
            if len(y) == 4 and len(m) == 2 and len(d) == 2:
                return f"{y}-{m}-{d}"

    # Formato ISO AAAA-MM-DD
    if "-" in s:
        partes = s.split("-")
        if len(partes) == 3 and len(partes[0]) == 4:
            return f"{partes[0]}-{partes[1].zfill(2)}-{partes[2].zfill(2)}"

    return data_fallback


def _calcular_hash_arquivo(caminho: Path) -> str:
    """Calcula o hash SHA-256 de um arquivo para prevenir reprocessamento desnecessário.

    Lê o arquivo em blocos de 64KB para manter uso de memória mínimo.

    Args:
        caminho: Caminho absoluto para o arquivo no disco.

    Returns:
        str: Hash hexadecimal SHA-256 ou string vazia em caso de falha de I/O.

    Complexity:
        Tempo: O(B) onde B é o tamanho do arquivo em bytes.
        Espaço: O(1) - buffer fixo de 64KB.
    """
    hasher = hashlib.sha256()
    try:
        with open(caminho, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                hasher.update(chunk)
        return hasher.hexdigest()
    except Exception:
        return ""


def _garantir_tabelas_e_indices(conn_triagem) -> None:
    """Garante tabelas de controle, índices e integridade para importação.

    Cria:
    - `historico_arquivos_importados`: Registro de auditoria com hash e contadores.
    - Índice único `uq_hist_cnpj_fiscal` em `historico_fiscais(cnpj_limpo, fiscal)`.
    - `controle_competencias`: Controle de meses ativos e arquivados.

    Args:
        conn_triagem: Conexão SQLite ativa com `triagem_vre.db`.

    Complexity:
        Tempo: O(1).
        Espaço: O(1).
    """
    with db_transaction(conn_triagem) as cur:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS historico_arquivos_importados (
                nome_arquivo TEXT PRIMARY KEY,
                hash_arquivo TEXT,
                fiscais_importados INTEGER,
                ceps_importados INTEGER,
                protocolos_importados INTEGER DEFAULT 0,
                data_importacao DATETIME DEFAULT CURRENT_TIMESTAMP
            );
            """
        )
        # Migração defensiva: adiciona coluna se a tabela já existia de versão anterior
        cur.execute("PRAGMA table_info(historico_arquivos_importados);")
        colunas = [r[1] for r in cur.fetchall()]
        if "protocolos_importados" not in colunas:
            try:
                cur.execute("ALTER TABLE historico_arquivos_importados ADD COLUMN protocolos_importados INTEGER DEFAULT 0;")
            except Exception:
                pass

        cur.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS uq_hist_cnpj_fiscal
            ON historico_fiscais(cnpj_limpo, fiscal);
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS controle_competencias (
                mes_competencia TEXT PRIMARY KEY,
                arquivado INTEGER DEFAULT 0,
                atualizado_em DATETIME DEFAULT CURRENT_TIMESTAMP
            );
            """
        )


def _mapear_colunas_aba(ws) -> tuple[int, dict[str, int]]:
    """Localiza a linha de cabeçalho e mapeia as posições 1-based das colunas de interesse.

    Varre as primeiras 6 linhas da aba procurando termos indicativos das colunas oficiais
    do VRE e das planilhas de trabalho da SMAURB.

    Args:
        ws: Aba da planilha do OpenPyXL em modo read-only.

    Returns:
        tuple[int, dict[str, int]]: Tupla contendo o número da linha de cabeçalho (1-based)
                                    e o dicionário mapeando nome do campo -> índice da coluna.

    Complexity:
        Tempo: O(R * C) onde R <= 6 e C <= 60.
        Espaço: O(C) para o dicionário mapeado.
    """
    for r in range(1, min(7, ws.max_row + 1)):
        valores_linha = [
            str(ws.cell(row=r, column=c).value or "").strip().upper()
            for c in range(1, min(ws.max_column + 1, 60))
        ]

        mapa: dict[str, int] = {}
        for idx, val in enumerate(valores_linha, start=1):
            if not val:
                continue

            if "CNPJ" in val and "cnpj" not in mapa:
                mapa["cnpj"] = idx
            elif "FISCAL" in val and "fiscal" not in mapa:
                mapa["fiscal"] = idx
            elif ("ÁREA" in val or "AREA" == val) and "area" not in mapa:
                mapa["area"] = idx
            elif "BAIRRO" in val and "bairro" not in mapa:
                mapa["bairro"] = idx
            elif "CEP" == val and "cep" not in mapa:
                mapa["cep"] = idx
            elif ("PROTOCOLOLICENCA" in val or "LICENCA" in val or "LICENÇA" in val) and "protocolo_licenca" not in mapa:
                mapa["protocolo_licenca"] = idx
            elif ("PROTOCOLOSIL" in val or "PROTOCOLO SIL" in val or val == "SIL") and "protocolo_sil" not in mapa:
                mapa["protocolo_sil"] = idx
            elif ("NOMEEMPRESARIAL" in val or "NOME EMPRESARIAL" in val or "RAZAO" in val or "RAZÃO" in val) and "nome_empresarial" not in mapa:
                mapa["nome_empresarial"] = idx
            elif "STATUS" in val and "status" not in mapa:
                mapa["status"] = idx
            elif ("TIPO DE SERVI" in val or "TIPO_SERVICO" in val or "SERVICO" in val or "SERVIÇO" in val) and "tipo_servico" not in mapa:
                mapa["tipo_servico"] = idx
            elif ("CNAESLICENCIADAS" in val or "CNAES" in val or "CNAE" in val) and "cnaes" not in mapa:
                mapa["cnaes"] = idx
            elif "CNAESNAOLICENCIADAS" in val and "cnaes_nao" not in mapa:
                mapa["cnaes_nao"] = idx
            elif ("INSCRICAOMUNICIPAL" in val or "INSCRIÇÃO" in val or "INSCRICAO" in val) and "inscricao_municipal" not in mapa:
                mapa["inscricao_municipal"] = idx
            elif ("GRAURISCOGERAL" in val or "RISCO" in val) and "grau_risco" not in mapa:
                mapa["grau_risco"] = idx
            elif ("SITUACAOPOSTURA" in val or "POSTURA" in val) and "situacao_postura" not in mapa:
                mapa["situacao_postura"] = idx
            elif "TELEFONE" in val and "telefone" not in mapa:
                mapa["telefone"] = idx
            elif ("E-MAIL" in val or "EMAIL" in val) and "email" not in mapa:
                mapa["email"] = idx
            elif ("TIPOENDERECO" in val or "TIPO ENDERECO" in val or "TIPO ENDEREÇO" in val) and "tipo_end" not in mapa:
                mapa["tipo_end"] = idx
            elif ("LOGRADOURO" in val or "RUA" in val or "AVENIDA" in val) and "logradouro" not in mapa:
                mapa["logradouro"] = idx
            elif ("NUMERO" in val or "NÚMERO" in val) and "numero" not in mapa:
                mapa["numero"] = idx
            elif "COMPLEMENTO" in val and "complemento" not in mapa:
                mapa["complemento"] = idx
            elif "FINALIDADEENDERECO" in val and "finalidade_end" not in mapa:
                mapa["finalidade_end"] = idx
            elif "IPTU" in val and "iptu" not in mapa:
                mapa["iptu"] = idx
            elif "PROTOCOLO EXISTENTE" in val and "protocolo_existente" not in mapa:
                mapa["protocolo_existente"] = idx
            elif ("PROTOCOLOSIGM" in val or "SIGM" in val) and "protocolo_sigm" not in mapa:
                mapa["protocolo_sigm"] = idx
            elif ("RESTRIÇÃO" == val or "RESTRICAO" == val) and "restricao" not in mapa:
                mapa["restricao"] = idx
            elif ("ENDEREÇO" == val or "ENDERECO" == val) and "endereco" not in mapa:
                mapa["endereco"] = idx
            elif "DATA" in val:
                if "DATA VRE" in val or "DATA DO VRE" in val:
                    mapa["data_sec"] = idx
                elif "data" not in mapa:
                    mapa["data"] = idx
                elif "data_sec" not in mapa:
                    mapa["data_sec"] = idx

        if "cnpj" in mapa and ("fiscal" in mapa or "protocolo_licenca" in mapa or "protocolo_sil" in mapa):
            return r, mapa

    return 0, {}


def reconciliar_fiscais_backlog(triagem_conn) -> int:
    """Atribui automaticamente o fiscal histórico mais recente para protocolos sem fiscal.

    Regra de Negócio:
    - Protocolos pendentes ou sem fiscal cujo CNPJ já possua atendimento anterior
      em `historico_fiscais` recebem o fiscal titular da empresa.
    - Se o protocolo possuir atividade noturna (`flag_noturno = 1`), o tipo de serviço
      é ajustado para 'Noturno'.
    - Se não for noturno, é classificado como 'Diurno Livre'.

    Args:
        triagem_conn: Conexão ativa com o banco central `triagem_vre.db`.

    Returns:
        int: Número de protocolos atualizados com fiscal histórico.

    Complexity:
        Tempo: O(P) onde P é o número de protocolos sem fiscal atribuído.
        Espaço: O(1).
    """
    with db_transaction(triagem_conn) as cur:
        cur.execute(
            """
            UPDATE vre_backlog
            SET fiscal_atribuido = (
                SELECT h.fiscal
                FROM historico_fiscais h
                WHERE h.cnpj_limpo = vre_backlog.cnpj_limpo
                ORDER BY h.data_atendimento DESC, h.id DESC
                LIMIT 1
            ),
            tipo_servico = CASE
                WHEN flag_noturno = 1 AND (tipo_servico = 'Noturno Livre' OR tipo_servico IS NULL OR TRIM(tipo_servico) = '')
                THEN 'Noturno'
                WHEN (flag_noturno = 0 OR flag_noturno IS NULL)
                     AND (tipo_servico IN ('Com Restrição Livre', 'Sem Restrição Livre', 'Com Restricao Livre', 'Sem Restricao Livre') OR tipo_servico IS NULL OR TRIM(tipo_servico) = '')
                THEN 'Diurno Livre'
                ELSE tipo_servico
            END
            WHERE (fiscal_atribuido IS NULL OR TRIM(fiscal_atribuido) = '' OR fiscal_atribuido = 'Não Informado')
              AND EXISTS (
                  SELECT 1 FROM historico_fiscais h WHERE h.cnpj_limpo = vre_backlog.cnpj_limpo
              );
            """
        )
        return cur.rowcount


def importar_planilha_historico(
    caminho_xlsx: Path | str,
    triagem_db_path: Path | str | None = None,
    cep_db_path: Path | str | None = None,
    forcar: bool = False,
) -> dict[str, Any]:
    """Importa histórico de fiscais, correlações de CEP e protocolos no backlog com deduplicação.

    Pipeline Completo:
    1. Verificação de integridade e cálculo de hash SHA-256 do arquivo.
    2. Checagem em `historico_arquivos_importados` para evitar reprocessamento idêntico (salvo se `forcar=True`).
    3. Varredura com OpenPyXL em modo `read_only=True` e `data_only=True` para máxima velocidade.
    4. Identificação dinâmica de colunas por aba.
    5. Deduplicação em memória retendo as datas mais recentes e descartando fiscais inválidos.
    6. Inserção atômica via UPSERT em `historico_fiscais`, `cep_mapeamento` e `vre_backlog`.
    7. Limpeza preventiva de indeferidos e reconciliação geográfica de áreas.

    Args:
        caminho_xlsx: Caminho para a planilha Excel a ser importada.
        triagem_db_path: Caminho alternativo do banco central.
        cep_db_path: Caminho alternativo do banco de CEPs.
        forcar: Se True, ignora a checagem de hash e reprocessa a planilha integralmente.

    Returns:
        dict[str, Any]: Dicionário com estatísticas completas da importação.

    Raises:
        FileNotFoundError: Se o arquivo especificado não existir.

    Complexity:
        Tempo: O(L) onde L é o número de linhas da planilha.
        Espaço: O(U) onde U é o número de registros únicos consolidados em memória.
    """
    p_arq = Path(caminho_xlsx).resolve()
    if not p_arq.exists() or not p_arq.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {caminho_xlsx}")

    conn_triagem = get_triagem_db(triagem_db_path)
    initialize_schema_if_needed(conn_triagem, is_cep=False)
    _garantir_tabelas_e_indices(conn_triagem)

    hash_atual = _calcular_hash_arquivo(p_arq)

    if not forcar and hash_atual:
        cur = conn_triagem.cursor()
        cur.execute(
            "SELECT fiscais_importados, ceps_importados, protocolos_importados FROM historico_arquivos_importados WHERE nome_arquivo = ? AND hash_arquivo = ?;",
            (p_arq.name, hash_atual),
        )
        row = cur.fetchone()
        if row:
            conn_triagem.close()
            return {
                "arquivo": p_arq.name,
                "status": "JA_IMPORTADO",
                "linhas_lidas": 0,
                "duplicatas_descartadas": 0,
                "fiscais_inseridos": row[0],
                "ceps_inseridos": row[1],
                "protocolos_inseridos": row[2] if len(row) > 2 else 0,
            }

    conn_cep = get_cep_db(cep_db_path)
    initialize_schema_if_needed(conn_cep, is_cep=True)

    try:
        wb = openpyxl.load_workbook(str(p_arq), data_only=True, read_only=True)
    except Exception as e:
        conn_triagem.close()
        conn_cep.close()
        return {
            "arquivo": p_arq.name,
            "status": "ERRO_LEITURA",
            "mensagem": str(e),
            "linhas_lidas": 0,
            "duplicatas_descartadas": 0,
            "fiscais_inseridos": 0,
            "ceps_inseridos": 0,
            "protocolos_inseridos": 0,
        }

    # Data fallback baseada na última modificação do arquivo
    data_mod_arquivo = datetime.date.fromtimestamp(p_arq.stat().st_mtime).isoformat()

    # Dicionários de deduplicação em memória
    # Chave fiscal: (cnpj_limpo, fiscal_limpo) -> (data_str_mais_recente, origem_str)
    fiscais_consolidados: dict[tuple[str, str], tuple[str, str]] = {}
    # Chave CEP: cep_limpo -> (area_str, bairro_str)
    ceps_consolidados: dict[str, tuple[str, str]] = {}
    # Protocolos para inserção/upsert no vre_backlog
    protocolos_consolidados: dict[str, dict[str, Any]] = {}
    # Meses de competência identificados
    meses_consolidados: set[str] = set()

    total_linhas_lidas = 0
    total_linhas_brutas_fiscais = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        linha_cabecalho, col_map = _mapear_colunas_aba(ws)
        if not col_map or "cnpj" not in col_map:
            continue

        idx_cnpj = col_map["cnpj"]
        idx_fiscal = col_map.get("fiscal")
        idx_data = col_map.get("data")
        idx_data_sec = col_map.get("data_sec")
        idx_area = col_map.get("area")
        idx_bairro = col_map.get("bairro")
        idx_cep = col_map.get("cep")
        idx_licenca = col_map.get("protocolo_licenca")
        idx_sil = col_map.get("protocolo_sil")
        idx_nome = col_map.get("nome_empresarial")
        idx_status = col_map.get("status")
        idx_tipo_servico = col_map.get("tipo_servico")
        idx_restricao = col_map.get("restricao")
        idx_endereco = col_map.get("endereco")
        idx_cnaes = col_map.get("cnaes")
        idx_cnaes_nao = col_map.get("cnaes_nao")
        idx_insc = col_map.get("inscricao_municipal")
        idx_grau_risco = col_map.get("grau_risco")
        idx_postura = col_map.get("situacao_postura")
        idx_tel = col_map.get("telefone")
        idx_email = col_map.get("email")
        idx_tipo_end = col_map.get("tipo_end")
        idx_logr = col_map.get("logradouro")
        idx_num = col_map.get("numero")
        idx_comp = col_map.get("complemento")
        idx_finalidade = col_map.get("finalidade_end")
        idx_iptu = col_map.get("iptu")
        idx_prot_exist = col_map.get("protocolo_existente")
        idx_prot_sigm = col_map.get("protocolo_sigm")

        linha_num = linha_cabecalho
        for row in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
            linha_num += 1
            if not row or not any(row):
                continue

            total_linhas_lidas += 1

            cnpj_raw = row[idx_cnpj - 1] if idx_cnpj <= len(row) else None
            if not cnpj_raw:
                continue

            cnpj_str = str(cnpj_raw).strip()
            # Ignora cabeçalhos repetidos no meio da planilha
            if cnpj_str.upper() in ["CNPJ", "C.N.P.J.", "C.N.P.J", ""]:
                continue

            cnpj_limpo = normalizar_cnpj(cnpj_raw)
            if len(cnpj_limpo) != 14:
                continue

            fiscal_raw = row[idx_fiscal - 1] if idx_fiscal and idx_fiscal <= len(row) else None
            data_raw = row[idx_data - 1] if idx_data and idx_data <= len(row) else None
            data_sec_raw = row[idx_data_sec - 1] if idx_data_sec and idx_data_sec <= len(row) else None
            area_raw = row[idx_area - 1] if idx_area and idx_area <= len(row) else None
            bairro_raw = row[idx_bairro - 1] if idx_bairro and idx_bairro <= len(row) else None
            cep_raw = row[idx_cep - 1] if idx_cep and idx_cep <= len(row) else None

            # Resolução e normalização de data
            # DEVE-SE USAR A DATA DO VRE NÃO A COLOCADA:
            # Em planilhas com duas colunas de data (ex: 'Planilha de Trabalho atual.xlsx'),
            # a coluna secundária (próxima a ProtocoloLicenca/Sil) contém a Data oficial do VRE,
            # enquanto a primeira coluna é a data colocada na atribuição/distribuição.
            if idx_data_sec:
                data_candidata = data_sec_raw if (data_sec_raw is not None and str(data_sec_raw).strip() not in ("", "None")) else data_raw
            else:
                data_candidata = data_raw

            data_normalizada = normalizar_data_atendimento(data_candidata, data_mod_arquivo)
            mes_competencia = data_normalizada[:7]
            if len(mes_competencia) == 7 and mes_competencia[4] == "-":
                meses_consolidados.add(mes_competencia)

            # 1. Higienização e Deduplicação do Fiscal
            fiscal_str = ""
            if fiscal_raw:
                f_limpo = str(fiscal_raw).strip()
                if f_limpo.upper() not in TERMOS_INVALIDOS_FISCAL and f_limpo not in (",", ";"):
                    fiscal_str = f_limpo

            if fiscal_str:
                total_linhas_brutas_fiscais += 1
                origem_nome = f"{p_arq.name}:{sheet_name}"
                chave_fiscal = (cnpj_limpo, fiscal_str)
                if chave_fiscal in fiscais_consolidados:
                    data_existente, _ = fiscais_consolidados[chave_fiscal]
                    if data_normalizada > data_existente:
                        fiscais_consolidados[chave_fiscal] = (data_normalizada, origem_nome)
                else:
                    fiscais_consolidados[chave_fiscal] = (data_normalizada, origem_nome)

            # 2. Higienização e Deduplicação do CEP x Área
            if cep_raw and area_raw:
                cep_str = str(cep_raw).strip()
                area_str = str(area_raw).strip().upper()

                if (
                    cep_str.upper() not in ["CEP", "C.E.P.", "C.E.P"]
                    and area_str not in ["ÁREA", "AREA", "ÁREA FISCAL", "AREA FISCAL"]
                    and area_str not in TERMOS_INVALIDOS_FISCAL
                ):
                    cep_limpo_calc = normalizar_cep(cep_raw)
                    if len(cep_limpo_calc) == 8:
                        bairro_limpo = normalizar_bairro(str(bairro_raw)) if bairro_raw else ""
                        if cep_limpo_calc not in ceps_consolidados:
                            ceps_consolidados[cep_limpo_calc] = (area_str, bairro_limpo)
                        elif not ceps_consolidados[cep_limpo_calc][1] and bairro_limpo:
                            ceps_consolidados[cep_limpo_calc] = (area_str, bairro_limpo)

            # 3. Consolidação de Protocolo para vre_backlog
            licenca_val = (
                str(row[idx_licenca - 1]).strip()
                if idx_licenca and idx_licenca <= len(row) and row[idx_licenca - 1]
                else ""
            )
            if not licenca_val or licenca_val.upper() in ["PROTOCOLOLICENCA", "PROTOCOLO LICENCA", "NONE", "NULL", "#N/A", "#REF!"]:
                licenca_val = f"HIST_{cnpj_limpo}_{data_normalizada}"

            sil_val = str(row[idx_sil - 1]).strip() if idx_sil and idx_sil <= len(row) and row[idx_sil - 1] else ""
            if sil_val.upper() in ["PROTOCOLOSIL", "PROTOCOLO SIL", "NONE", "NULL", "#N/A", "#REF!"]:
                sil_val = ""

            nome_val = str(row[idx_nome - 1]).strip() if idx_nome and idx_nome <= len(row) and row[idx_nome - 1] else ""
            if nome_val.upper() in ["NOMEEMPRESARIAL", "NOME EMPRESARIAL", "RAZAO SOCIAL", "NONE", "NULL"]:
                nome_val = ""

            bairro_val = str(row[idx_bairro - 1]).strip() if idx_bairro and idx_bairro <= len(row) and row[idx_bairro - 1] else ""
            area_val = str(row[idx_area - 1]).strip().upper() if idx_area and idx_area <= len(row) and row[idx_area - 1] else ""
            if area_val in ["ÁREA", "AREA", "NONE", "NULL"] or area_val in TERMOS_INVALIDOS_FISCAL:
                area_val = ""

            cnaes_val = str(row[idx_cnaes - 1]).strip() if idx_cnaes and idx_cnaes <= len(row) and row[idx_cnaes - 1] else ""
            tipo_servico_val = str(row[idx_tipo_servico - 1]).strip() if idx_tipo_servico and idx_tipo_servico <= len(row) and row[idx_tipo_servico - 1] else ""
            restricao_val = str(row[idx_restricao - 1]).strip() if idx_restricao and idx_restricao <= len(row) and row[idx_restricao - 1] else ""

            status_raw_orig = str(row[idx_status - 1]).strip() if idx_status and idx_status <= len(row) and row[idx_status - 1] else ""
            status_raw = status_raw_orig.upper()
            if status_raw in STATUS_CANONICOS:
                status_calc = STATUS_CANONICOS[status_raw]
                status_extra = STATUS_CANONICOS[status_raw]
            elif status_raw in ["FINALIZADO"]:
                status_calc = "FINALIZADO"
                status_extra = ""
            elif status_raw in ["DISTRIBUIDO", "DISTRIBUÍDO"]:
                status_calc = "DISTRIBUIDO"
                status_extra = ""
            elif status_raw in ["PENDENTE", "NÃO DISTRIBUÍDO", "NAO DISTRIBUIDO", "NÃO DISTRIBUIDO", "NAO DISTRIBUÍDO", "", "NONE", "NULL"]:
                status_calc = "DISTRIBUIDO" if fiscal_str else "PENDENTE"
                status_extra = ""
            else:
                status_calc = status_raw_orig
                status_extra = status_raw_orig

            postura_str = str(row[idx_postura - 1]).strip() if idx_postura and idx_postura <= len(row) and row[idx_postura - 1] else ""
            eh_indeferido_hist = (
                postura_str.upper() == "INDEFERIDO"
                or status_calc.upper() == "INDEFERIDO"
            )
            if eh_indeferido_hist:
                tipo_servico_val = ""
                restricao_val = ""

            comb_classificacao = f"{tipo_servico_val} {restricao_val}".upper()
            flag_noturno = 1 if "NOTURNO" in comb_classificacao else 0
            flag_restricao = 1 if (("RESTRIÇÃO" in comb_classificacao or "RESTRICAO" in comb_classificacao) and "SEM RESTRI" not in comb_classificacao) else 0

            iptu_str = str(row[idx_iptu - 1]).strip() if idx_iptu and idx_iptu <= len(row) and row[idx_iptu - 1] is not None else ""
            if iptu_str.endswith(".0") and iptu_str[:-2].isdigit():
                iptu_str = iptu_str[:-2]

            extras: dict[str, Any] = {
                "Iptu": iptu_str,
                "IPTU": iptu_str,
                "GrauRiscoGeral": str(row[idx_grau_risco - 1]).strip() if idx_grau_risco and idx_grau_risco <= len(row) and row[idx_grau_risco - 1] else "",
                "SituacaoPostura": postura_str,
                "InscricaoMunicipal": str(row[idx_insc - 1]).strip() if idx_insc and idx_insc <= len(row) and row[idx_insc - 1] else "",
                "Telefone": str(row[idx_tel - 1]).strip() if idx_tel and idx_tel <= len(row) and row[idx_tel - 1] else "",
                "E-mail": str(row[idx_email - 1]).strip() if idx_email and idx_email <= len(row) and row[idx_email - 1] else "",
                "TipoEndereco": str(row[idx_tipo_end - 1]).strip() if idx_tipo_end and idx_tipo_end <= len(row) and row[idx_tipo_end - 1] else "",
                "Logradouro": str(row[idx_logr - 1]).strip() if idx_logr and idx_logr <= len(row) and row[idx_logr - 1] else "",
                "Numero": str(row[idx_num - 1]).strip() if idx_num and idx_num <= len(row) and row[idx_num - 1] else "",
                "Complemento": str(row[idx_comp - 1]).strip() if idx_comp and idx_comp <= len(row) and row[idx_comp - 1] else "",
                "Bairro": bairro_val,
                "Cep": str(row[idx_cep - 1]).strip() if idx_cep and idx_cep <= len(row) and row[idx_cep - 1] is not None else "",
                "FinalidadeEndereco": str(row[idx_finalidade - 1]).strip() if idx_finalidade and idx_finalidade <= len(row) and row[idx_finalidade - 1] else "",
                "CnaesLicenciadas": cnaes_val,
                "CnaesNaoLicenciadas": str(row[idx_cnaes_nao - 1]).strip() if idx_cnaes_nao and idx_cnaes_nao <= len(row) and row[idx_cnaes_nao - 1] else "",
                "ProtocoloExistente": str(row[idx_prot_exist - 1]).strip() if idx_prot_exist and idx_prot_exist <= len(row) and row[idx_prot_exist - 1] else "",
                "ProtocoloSigm": str(row[idx_prot_sigm - 1]).strip() if idx_prot_sigm and idx_prot_sigm <= len(row) and row[idx_prot_sigm - 1] else "",
                "Data": data_normalizada,
            }
            if status_extra:
                extras["Status"] = status_extra
            if restricao_val and restricao_val.upper() not in ("NONE", "NULL"):
                extras["RestricaoCustom"] = restricao_val

            protocolos_consolidados[licenca_val] = {
                "protocolo_licenca": licenca_val,
                "protocolo_sil": sil_val,
                "cnpj_limpo": cnpj_limpo,
                "cnpj_formatado": formatar_cnpj(cnpj_limpo),
                "nome_empresarial": nome_val,
                "mes_competencia": mes_competencia,
                "data_importacao": data_normalizada,
                "bairro_original": bairro_val,
                "area": area_val,
                "cnaes_licenciadas": cnaes_val,
                "fiscal_atribuido": fiscal_str,
                "flag_restricao": flag_restricao,
                "flag_noturno": flag_noturno,
                "tipo_servico": tipo_servico_val,
                "status": status_calc,
                "dados_extras_json": json.dumps(extras, ensure_ascii=False),
            }

    wb.close()

    duplicatas_descartadas = total_linhas_brutas_fiscais - len(fiscais_consolidados)

    # 1. Inserção atômica com UPSERT no SQLite (historico_fiscais)
    if fiscais_consolidados:
        registros_upsert = [
            (cnpj, fiscal, data_atend, orig)
            for (cnpj, fiscal), (data_atend, orig) in fiscais_consolidados.items()
        ]
        with db_transaction(conn_triagem) as cur:
            cur.executemany(
                """
                INSERT INTO historico_fiscais (cnpj_limpo, fiscal, data_atendimento, origem_registro)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(cnpj_limpo, fiscal) DO UPDATE SET
                    data_atendimento = CASE
                        WHEN excluded.data_atendimento > historico_fiscais.data_atendimento THEN excluded.data_atendimento
                        ELSE historico_fiscais.data_atendimento
                    END,
                    origem_registro = excluded.origem_registro;
                """,
                registros_upsert,
            )

    # 2. Inserção atômica no banco de CEPs (cep_areas.db)
    for cep_limpo, (area_str, bairro_str) in ceps_consolidados.items():
        registrar_area_cep(
            cep=cep_limpo,
            area=area_str,
            bairro=bairro_str,
            origem=f"IMPORTACAO_{p_arq.name}",
            conn=conn_cep,
        )

    # 3. Inserção dos protocolos no vre_backlog
    if protocolos_consolidados:
        regs_backlog = [
            (
                p["protocolo_licenca"],
                p["protocolo_sil"],
                p["cnpj_limpo"],
                p["cnpj_formatado"],
                p["nome_empresarial"],
                p["mes_competencia"],
                p["data_importacao"],
                p["bairro_original"],
                p["area"],
                p["cnaes_licenciadas"],
                p["fiscal_atribuido"],
                p["flag_restricao"],
                p["flag_noturno"],
                p["tipo_servico"],
                p["status"],
                p["dados_extras_json"],
            )
            for p in protocolos_consolidados.values()
        ]
        with db_transaction(conn_triagem) as cur:
            cur.executemany(
                """
                INSERT INTO vre_backlog (
                    protocolo_licenca, protocolo_sil, cnpj_limpo, cnpj_formatado,
                    nome_empresarial, mes_competencia, data_importacao,
                    bairro_original, area, cnaes_licenciadas, fiscal_atribuido,
                    flag_restricao, flag_noturno, tipo_servico, status, dados_extras_json
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(protocolo_licenca) DO UPDATE SET
                    data_importacao = excluded.data_importacao,
                    mes_competencia = excluded.mes_competencia,
                    fiscal_atribuido = CASE
                        WHEN excluded.fiscal_atribuido IS NOT NULL AND TRIM(excluded.fiscal_atribuido) != ''
                        THEN excluded.fiscal_atribuido
                        ELSE vre_backlog.fiscal_atribuido
                    END,
                    area = CASE
                        WHEN excluded.area IS NOT NULL AND TRIM(excluded.area) != ''
                        THEN excluded.area
                        ELSE vre_backlog.area
                    END,
                    tipo_servico = CASE
                        WHEN excluded.tipo_servico IS NOT NULL AND TRIM(excluded.tipo_servico) != ''
                        THEN excluded.tipo_servico
                        ELSE vre_backlog.tipo_servico
                    END,
                    flag_restricao = CASE
                        WHEN excluded.flag_restricao = 1 THEN 1
                        ELSE vre_backlog.flag_restricao
                    END,
                    flag_noturno = CASE
                        WHEN excluded.flag_noturno = 1 THEN 1
                        ELSE vre_backlog.flag_noturno
                    END,
                    status = CASE
                        WHEN excluded.status IS NOT NULL AND TRIM(excluded.status) != '' AND excluded.status != 'PENDENTE'
                        THEN excluded.status
                        ELSE vre_backlog.status
                    END,
                    dados_extras_json = COALESCE(excluded.dados_extras_json, vre_backlog.dados_extras_json);
                """,
                regs_backlog,
            )

    # 4. Registro dos meses descobertos na tabela de controle de competências
    if meses_consolidados:
        with db_transaction(conn_triagem) as cur:
            for mes in meses_consolidados:
                if mes and len(mes) == 7:
                    cur.execute(
                        """
                        INSERT OR IGNORE INTO controle_competencias (mes_competencia, arquivado)
                        VALUES (?, 0);
                        """,
                        (mes,),
                    )

    # 5. Registra o arquivo como importado
    with db_transaction(conn_triagem) as cur:
        cur.execute(
            """
            INSERT OR REPLACE INTO historico_arquivos_importados
            (nome_arquivo, hash_arquivo, fiscais_importados, ceps_importados, protocolos_importados, data_importacao)
            VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP);
            """,
            (
                p_arq.name,
                hash_atual,
                len(fiscais_consolidados),
                len(ceps_consolidados),
                len(protocolos_consolidados),
            ),
        )

    # 6. Reconcilia o backlog ativo (preenche fiscais e áreas para protocolos sem preenchimento e limpa indeferidos)
    reconciliar_areas_backlog(triagem_conn=conn_triagem, cep_conn=conn_cep)
    reconciliar_fiscais_backlog(triagem_conn=conn_triagem)
    from src.core.excel_writer import limpar_indeferidos_e_normalizar_backlog
    limpar_indeferidos_e_normalizar_backlog(conn_triagem, executar_migracao_col_f=False)

    conn_triagem.close()
    conn_cep.close()

    return {
        "arquivo": p_arq.name,
        "status": "SUCESSO",
        "linhas_lidas": total_linhas_lidas,
        "duplicatas_descartadas": max(0, duplicatas_descartadas),
        "fiscais_inseridos": len(fiscais_consolidados),
        "ceps_inseridos": len(ceps_consolidados),
        "protocolos_inseridos": len(protocolos_consolidados),
        "meses_identificados": sorted(list(meses_consolidados)),
    }


def escanear_e_importar_planilhas_antigas(
    pasta_origem: Path | str | None = None,
    forcar: bool = False,
) -> dict[str, Any]:
    """Varre 'Planilhas anteriores/' e a raiz em busca de planilhas antigas e importa tudo com deduplicação.

    Ordem de Processamento:
    1. Localiza planilhas em `Planilhas anteriores/`, na raiz e em `Documentação do projeto/Documentos de exemplo`.
    2. Ordena os arquivos pela data de modificação (`st_mtime`) em ordem cronológica crescente.
       Dessa forma, planilhas mais recentes sobrescrevem planilhas mais antigas durante os UPSERTs.
    3. Retorna sumário consolidado com contadores agregados e detalhes de cada arquivo processado.

    Args:
        pasta_origem: Pasta customizada para varredura ou None para usar `PLANILHAS_ANTERIORES_DIR`.
        forcar: Se True, força a reimportação mesmo se o hash do arquivo já foi registrado.

    Returns:
        dict[str, Any]: Estatísticas completas da carga em lote.

    Complexity:
        Tempo: O(F * L) onde F é a quantidade de planilhas e L a média de linhas por planilha.
        Espaço: O(U) proporcional ao total consolidado de entidades únicas.
    """
    dir_busca = Path(pasta_origem) if pasta_origem else PLANILHAS_ANTERIORES_DIR

    arquivos_candidatos: list[Path] = []

    # 1. Arquivos em 'Planilhas anteriores/'
    if dir_busca.exists() and dir_busca.is_dir():
        for ext in ["*.xlsx", "*.xls"]:
            arquivos_candidatos.extend(dir_busca.glob(ext))

    # 2. Arquivos na raiz do BASE_DIR que não sejam a planilha ativa ou templates
    if BASE_DIR.exists() and BASE_DIR.is_dir():
        for ext in ["*.xlsx", "*.xls"]:
            for f in BASE_DIR.glob(ext):
                nome_upper = f.name.upper()
                if "ATIVA" in nome_upper or "TEMPLATE" in nome_upper:
                    continue
                if f not in arquivos_candidatos:
                    arquivos_candidatos.append(f)

    # 3. Documentação do projeto / Documentos de exemplo
    dir_docs = BASE_DIR / "Documentação do projeto" / "Documentos de exemplo"
    if dir_docs.exists() and dir_docs.is_dir():
        for ext in ["*.xlsx", "*.xls"]:
            for f in dir_docs.glob(ext):
                nome_upper = f.name.upper()
                if "CNAI" in nome_upper or "BAIRRO" in nome_upper:
                    continue
                if f not in arquivos_candidatos:
                    arquivos_candidatos.append(f)

    if not arquivos_candidatos:
        return {
            "total_arquivos": 0,
            "linhas_totais": 0,
            "total_duplicatas_descartadas": 0,
            "total_fiscais": 0,
            "total_ceps": 0,
            "total_protocolos": 0,
            "meses": [],
            "detalhes": [],
        }

    # Ordena do mais antigo para o mais recente para que edições recentes prevaleçam no UPSERT
    arquivos_candidatos.sort(key=lambda f: f.stat().st_mtime if f.exists() else 0)

    linhas_totais = 0
    total_duplicatas = 0
    total_fiscais = 0
    total_ceps = 0
    total_protocolos = 0
    todos_meses: set[str] = set()
    detalhes: list[dict[str, Any]] = []

    for arq in arquivos_candidatos:
        res = importar_planilha_historico(arq, forcar=forcar)
        linhas_totais += res.get("linhas_lidas", 0)
        total_duplicatas += res.get("duplicatas_descartadas", 0)
        total_fiscais += res.get("fiscais_inseridos", 0)
        total_ceps += res.get("ceps_inseridos", 0)
        total_protocolos += res.get("protocolos_inseridos", 0)
        for m in res.get("meses_identificados", []):
            todos_meses.add(m)
        detalhes.append(res)

    return {
        "total_arquivos": len(arquivos_candidatos),
        "linhas_totais": linhas_totais,
        "total_duplicatas_descartadas": total_duplicatas,
        "total_fiscais": total_fiscais,
        "total_ceps": total_ceps,
        "total_protocolos": total_protocolos,
        "meses": sorted(list(todos_meses)),
        "detalhes": detalhes,
    }
