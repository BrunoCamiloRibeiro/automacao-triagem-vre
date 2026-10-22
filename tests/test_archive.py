"""Testes para o serviço de arquivamento e visibilidade de competências por mês."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.admin.archive_service import (
    alternar_status_competencia,
    apagar_dia_banco,
    aplicar_reorganizacao_planilha,
    arquivar_competencia,
    buscar_historico_fiscais,
    definir_status_competencia,
    listar_competencias_status,
    listar_dias_importados,
    normalizar_data_filtro,
    obter_estatisticas_gerais,
    reativar_competencia,
)
from src.core.database import connect_db, initialize_schema_if_needed


class TestArchiveService(unittest.TestCase):
    """Verifica a gestão de visibilidade por mês, backup da planilha e integridade do backlog."""

    def setUp(self) -> None:
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmp_dir.name) / "test_archive.db"
        self.conn = connect_db(self.db_path)
        initialize_schema_if_needed(self.conn, is_cep=False)

        # Planilha de trabalho temporária
        self.saida_xlsx = Path(self.tmp_dir.name) / "Planilha_Controle_VRE_Ativa.xlsx"

    def tearDown(self) -> None:
        self.conn.close()
        self.tmp_dir.cleanup()

    def test_listar_competencias_e_status(self) -> None:
        """Verifica se todos os meses registrados aparecem com seu status de arquivamento."""
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, cnpj_limpo, mes_competencia, data_importacao, status, fiscal_atribuido)
            VALUES ('P1', '11111111000111', '2025-08', '2025-08-01', 'FINALIZADO', 'Fiscal_Alfa'),
                   ('P2', '22222222000122', '2025-08', '2025-08-01', 'PENDENTE', NULL),
                   ('P3', '33333333000133', '2025-09', '2025-09-01', 'DISTRIBUIDO', 'Fiscal_Beta');
            """
        )

        lotes = listar_competencias_status(self.conn)
        self.assertEqual(len(lotes), 2)
        meses = [l["mes_competencia"] for l in lotes]
        self.assertIn("2025-08", meses)
        self.assertIn("2025-09", meses)

        # Por padrão, meses entram como ativos (arquivado = False)
        for l in lotes:
            self.assertFalse(l["arquivado"])
            self.assertIn("Ativo", l["status_texto"])

    def test_alternar_e_definir_status_competencia(self) -> None:
        """Verifica a alternância entre Ativo e Arquivado (Oculto)."""
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, cnpj_limpo, mes_competencia, data_importacao, status)
            VALUES ('P1', '11111111000111', '2025-08', '2025-08-01', 'PENDENTE');
            """
        )

        # Arquiva a competência
        ok = definir_status_competencia("2025-08", arquivado=True, conn=self.conn)
        self.assertTrue(ok)

        lotes = listar_competencias_status(self.conn)
        self.assertTrue(lotes[0]["arquivado"])
        self.assertIn("Arquivado", lotes[0]["status_texto"])

        # Alterna para reativar
        novo_status = alternar_status_competencia("2025-08", conn=self.conn)
        self.assertFalse(novo_status)  # Agora arquivado = False (Ativo)

        lotes_pos = listar_competencias_status(self.conn)
        self.assertFalse(lotes_pos[0]["arquivado"])

        # Reativa explicitamente
        reativar_competencia("2025-08", conn=self.conn)
        stats = obter_estatisticas_gerais(self.conn)
        self.assertEqual(stats["meses_ativos"], 1)
        self.assertEqual(stats["meses_arquivados"], 0)

    def test_reorganizacao_planilha_filtra_arquivados_e_faz_backup(self) -> None:
        """Testa o fluxo completo: salva planilha anterior como backup,

        filtra mês arquivado e regenera planilha com apenas os meses ativos,
        mantendo os dados preservados no banco SQLite.
        """
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, protocolo_sil, cnpj_limpo, cnpj_formatado, nome_empresarial, mes_competencia, data_importacao, status, flag_restricao, flag_noturno, tipo_servico)
            VALUES ('P10', 'SIL-10', '11111111000111', '11.111.111/0001-11', 'EMPRESA MES 8', '2025-08', '2025-08-01', 'PENDENTE', 0, 0, 'Sem Restricao Livre'),
                   ('P20', 'SIL-20', '22222222000122', '22.222.222/0001-22', 'EMPRESA MES 9', '2025-09', '2025-09-01', 'PENDENTE', 0, 0, 'Sem Restricao Livre');
            """
        )

        # 1. Cria uma planilha ativa inicial simulando a versão em uso
        wb_inicial = openpyxl.Workbook()
        ws_ini = wb_inicial.active
        ws_ini.cell(1, 1, "VERSAO ANTERIOR")
        wb_inicial.save(self.saida_xlsx)
        wb_inicial.close()
        self.assertTrue(self.saida_xlsx.exists())

        # 2. Arquiva o mês 2025-08 (ocultando-o da planilha)
        definir_status_competencia("2025-08", arquivado=True, conn=self.conn)
        definir_status_competencia("2025-09", arquivado=False, conn=self.conn)

        # 3. Executa a reorganização
        res = aplicar_reorganizacao_planilha(conn=self.conn, saida_xlsx=self.saida_xlsx)
        self.assertEqual(res["status"], "SUCESSO")
        self.assertEqual(res["linhas_geradas"], 1)
        self.assertIn("2025-09", res["meses_ativos"])
        self.assertIn("2025-08", res["meses_arquivados"])

        # 4. Verifica a planilha gerada: deve conter apenas a empresa do mês 9
        wb_novo = openpyxl.load_workbook(self.saida_xlsx, data_only=True)
        ws_novo = wb_novo.active
        # Linha 1: Meta, Linha 2: Header, Linha 3: Dado
        self.assertEqual(ws_novo.max_row, 3)
        self.assertEqual(ws_novo.cell(3, 19).value, "EMPRESA MES 9")
        wb_novo.close()

        # 5. Verifica integridade do banco: ambos os protocolos continuam no SQLite!
        cur.execute("SELECT COUNT(*) FROM vre_backlog;")
        self.assertEqual(cur.fetchone()[0], 2)

        # 6. Reativa o mês 2025-08 e aplica reorganização novamente
        definir_status_competencia("2025-08", arquivado=False, conn=self.conn)
        res2 = aplicar_reorganizacao_planilha(conn=self.conn, saida_xlsx=self.saida_xlsx)
        self.assertEqual(res2["linhas_geradas"], 2)

        wb_ambos = openpyxl.load_workbook(self.saida_xlsx, data_only=True)
        ws_ambos = wb_ambos.active
        self.assertEqual(ws_ambos.max_row, 4)  # Linha 1 + Cabeçalho + 2 dados
        wb_ambos.close()

    def test_busca_historico_fiscais(self) -> None:
        """Verifica busca de fiscais por CNPJ e Razão Social."""
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO historico_fiscais (cnpj_limpo, fiscal, data_atendimento, origem_registro)
            VALUES ('33333333000133', 'Fiscal_Alfa', '2025-07-15', 'Importacao_Manual');
            """
        )
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, cnpj_limpo, mes_competencia, data_importacao, fiscal_atribuido, nome_empresarial)
            VALUES ('P99', '44444444000144', '2025-09', '2025-09-01', 'Fiscal_Beta', 'PADARIA BOA SORTE LTDA');
            """
        )

        # Busca por CNPJ
        res1 = buscar_historico_fiscais("33.333.333/0001-33", conn=self.conn)
        self.assertEqual(len(res1), 1)
        self.assertEqual(res1[0]["fiscal"], "Fiscal_Alfa")

        # Busca por Razão Social
        res2 = buscar_historico_fiscais("BOA SORTE", conn=self.conn)
        self.assertEqual(len(res2), 1)
        self.assertEqual(res2[0]["fiscal"], "Fiscal_Beta")

    def test_normalizar_data_filtro(self) -> None:
        """Verifica a harmonização de formatos heterogêneos de data para AAAA-MM-DD."""
        self.assertEqual(normalizar_data_filtro("21/10/2026"), "2026-10-21")
        self.assertEqual(normalizar_data_filtro("2026-10-21"), "2026-10-21")
        self.assertEqual(normalizar_data_filtro("21-10-2026"), "2026-10-21")
        self.assertEqual(normalizar_data_filtro("2026/10/21"), "2026-10-21")
        self.assertEqual(normalizar_data_filtro(""), "")
        self.assertEqual(normalizar_data_filtro(None), "")

    def test_listar_dias_importados(self) -> None:
        """Verifica a consolidação correta de dias, totais e status por remessa diária."""
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, cnpj_limpo, mes_competencia, data_importacao, status)
            VALUES ('D1_1', '11111111000111', '2026-10', '2026-10-20', 'PENDENTE'),
                   ('D1_2', '22222222000122', '2026-10', '2026-10-20', 'FINALIZADO'),
                   ('D2_1', '33333333000133', '2026-10', '2026-10-21', 'DISTRIBUIDO');
            """
        )

        dias = listar_dias_importados(self.conn)
        self.assertEqual(len(dias), 2)

        # Ordenado por data mais recente primeiro
        self.assertEqual(dias[0]["data_importacao"], "2026-10-21")
        self.assertEqual(dias[0]["data_formatada"], "21/10/2026")
        self.assertEqual(dias[0]["total"], 1)
        self.assertEqual(dias[0]["distribuidos"], 1)

        self.assertEqual(dias[1]["data_importacao"], "2026-10-20")
        self.assertEqual(dias[1]["data_formatada"], "20/10/2026")
        self.assertEqual(dias[1]["total"], 2)
        self.assertEqual(dias[1]["pendentes"], 1)
        self.assertEqual(dias[1]["finalizados"], 1)

    def test_apagar_dia_banco(self) -> None:
        """Verifica a exclusão atômica de uma remessa diária específica e regeração da planilha."""
        cur = self.conn.cursor()
        cur.execute(
            """
            INSERT INTO vre_backlog (protocolo_licenca, protocolo_sil, cnpj_limpo, cnpj_formatado, nome_empresarial, mes_competencia, data_importacao, status, flag_restricao, flag_noturno, tipo_servico)
            VALUES ('P_ERR_1', 'SIL-1', '11111111000111', '11.111.111/0001-11', 'EMPRESA ERRADA 1', '2026-10', '2026-10-21', 'PENDENTE', 0, 0, 'Sem Restricao Livre'),
                   ('P_ERR_2', 'SIL-2', '22222222000122', '22.222.222/0001-22', 'EMPRESA ERRADA 2', '2026-10', '2026-10-21', 'PENDENTE', 0, 0, 'Sem Restricao Livre'),
                   ('P_OK_1',  'SIL-3', '33333333000133', '33.333.333/0001-33', 'EMPRESA CORRETA',  '2026-10', '2026-10-20', 'PENDENTE', 0, 0, 'Sem Restricao Livre');
            """
        )

        # 1. Cria planilha ativa inicial com os 3 registros
        res_ini = aplicar_reorganizacao_planilha(conn=self.conn, saida_xlsx=self.saida_xlsx)
        self.assertEqual(res_ini["linhas_geradas"], 3)

        # 2. Apaga o dia errado informado no formato brasileiro '21/10/2026'
        res_del = apagar_dia_banco(
            "21/10/2026",
            conn=self.conn,
            atualizar_planilha=True,
            saida_xlsx=self.saida_xlsx,
        )
        self.assertTrue(res_del["sucesso"])
        self.assertEqual(res_del["removidos"], 2)
        self.assertEqual(res_del["data_iso"], "2026-10-21")
        self.assertEqual(res_del["linhas_planilha_ativa"], 1)

        # 3. Verifica integridade do SQLite: apenas o protocolo de 2026-10-20 permanece
        cur.execute("SELECT protocolo_licenca FROM vre_backlog;")
        restantes = [r[0] for r in cur.fetchall()]
        self.assertEqual(restantes, ["P_OK_1"])

        # 4. Verifica se a planilha ativa no disco foi atualizada e tem apenas 1 linha de dados
        wb = openpyxl.load_workbook(self.saida_xlsx, data_only=True)
        ws = wb.active
        # Linha 1: Meta, Linha 2: Header, Linha 3: 'EMPRESA CORRETA'
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(3, 19).value, "EMPRESA CORRETA")
        wb.close()

        # 5. Tenta apagar data inexistente: deve retornar sucesso=False e removidos=0
        res_inex = apagar_dia_banco("15/01/2020", conn=self.conn, atualizar_planilha=False)
        self.assertFalse(res_inex["sucesso"])
        self.assertEqual(res_inex["removidos"], 0)


if __name__ == "__main__":
    unittest.main()
