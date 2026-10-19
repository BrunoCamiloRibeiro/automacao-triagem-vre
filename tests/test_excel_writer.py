"""Testes para o módulo excel_writer.py."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.core.excel_writer import criar_template_padrao, gerar_planilha_trabalho


class TestExcelWriter(unittest.TestCase):
    """Verifica geração atômica de Excel, Plain Values, congelamento de painel e estilos."""

    def setUp(self) -> None:
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.template_path = Path(self.tmp_dir.name) / "test_template.xlsx"
        self.output_path = Path(self.tmp_dir.name) / "test_output.xlsx"
        criar_template_padrao(self.template_path)

    def tearDown(self) -> None:
        self.tmp_dir.cleanup()

    def test_estrutura_template(self) -> None:
        wb = openpyxl.load_workbook(str(self.template_path))
        self.assertIn("Controle VRE", wb.sheetnames)
        ws = wb["Controle VRE"]

        self.assertEqual(ws.freeze_panes, "A3")
        self.assertTrue(ws.auto_filter.ref.startswith("A2:AF"))
        self.assertEqual(ws.cell(row=1, column=23).value, "SÃO JOSÉ DO RIO PRETO")
        self.assertEqual(ws.cell(row=2, column=1).value, "Número ordem ")
        self.assertEqual(ws.cell(row=2, column=2).value, "Status")
        self.assertEqual(ws.cell(row=2, column=3).value, "Fiscal")
        self.assertEqual(ws.cell(row=2, column=4).value, "Tipo de Serviço")
        self.assertEqual(ws.cell(row=2, column=5).value, "Restrição")
        self.assertEqual(ws.cell(row=2, column=6).value, "Data")
        self.assertEqual(ws.cell(row=2, column=7).value, "IPTU")
        self.assertEqual(ws.cell(row=2, column=8).value, "ÁREA")
        self.assertEqual(ws.cell(row=2, column=9).value, "Endereço")
        self.assertEqual(ws.cell(row=2, column=32).value, "CEP4")
        self.assertEqual(ws.max_column, 32)

        # Valida as regras de cores na ordem institucional (incluindo Sem Restrição amarelo FFFFC000)
        self.assertEqual(len(ws.conditional_formatting), 1)
        cf_block = list(ws.conditional_formatting)[0]
        self.assertEqual(len(cf_block.rules), 15)
        self.assertEqual(cf_block.rules[0].formula[0], '$B3="Em Andamento"')
        self.assertEqual(cf_block.rules[5].formula[0], '$P3="INDEFERIDO"')
        self.assertEqual(cf_block.rules[6].formula[0], '$P3="CANCELADO"')
        self.assertEqual(cf_block.rules[7].formula[0], '$D3="Noturno entregue"')
        self.assertEqual(cf_block.rules[8].formula[0], '$C3<>""')
        self.assertEqual(cf_block.rules[9].formula[0], '$D3="Noturno Livre"')
        self.assertEqual(cf_block.rules[10].formula[0], '$E3="Noturno"')


    def test_gerar_planilha_plain_values_e_flags(self) -> None:
        dados_mock = [
            {
                "protocolo_licenca": "PROT_001",
                "protocolo_sil": "SIL_101",
                "cnpj_limpo": "61017032000194",
                "cnpj_formatado": "61.017.032/0001-94",
                "nome_empresarial": "EMPRESA NOTURNA LTDA",
                "status": "PENDENTE",
                "fiscal_atribuido": "Fiscal_Alfa",
                "tipo_servico": "Noturno Livre",
                "flag_restricao": 0,
                "flag_noturno": 1,
                "area": "CENTRAL",
                "bairro_original": "VILA ANCHIETA",
                "dados_extras_json": json.dumps(
                    {
                        "TipoEndereco": "RUA",
                        "Logradouro": "AIMORES",
                        "Numero": "287",
                        "Complemento": ":",
                        "Cep": "15050090",
                        "Iptu": "0307968000",
                        "Data": "15/05/2026",
                    }
                ),
            },
            {
                "protocolo_licenca": "PROT_002",
                "protocolo_sil": "SIL_102",
                "cnpj_limpo": "07244271000131",
                "cnpj_formatado": "07.244.271/0001-31",
                "nome_empresarial": "EMPRESA COMPLEXA LTDA",
                "status": "DISTRIBUIDO",
                "fiscal_atribuido": "Fiscal_Beta",
                "tipo_servico": "Com Restrição Livre",
                "flag_restricao": 1,
                "flag_noturno": 0,
                "area": "HB",
                "bairro_original": "JARDIM SÃO MARCO",
                "dados_extras_json": json.dumps(
                    {
                        "TipoEndereco": "AV",
                        "Logradouro": "MARIO ANDREAZZA",
                        "Numero": "181",
                        "Complemento": "SALA 1",
                        "Cep": "15081490",
                        "Iptu": "0554051005",
                        "Data": "16/05/2026",
                    }
                ),
            },
        ]

        gerar_planilha_trabalho(
            registros_backlog=dados_mock,
            caminho_saida=self.output_path,
            caminho_template=self.template_path,
        )

        wb_out = openpyxl.load_workbook(str(self.output_path))
        ws = wb_out["Controle VRE"]

        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.max_column, 32)
        self.assertEqual(ws.freeze_panes, "A3")

        # Linha 3 (Noturno ativo, status PENDENTE deve ficar vazio)
        self.assertEqual(ws.cell(row=3, column=1).value, 1)  # Seq 1
        self.assertIn(ws.cell(row=3, column=2).value, ("", None)) # Status vazio (não distribuído / pendente)
        self.assertEqual(ws.cell(row=3, column=3).value, "Fiscal_Alfa")
        self.assertEqual(ws.cell(row=3, column=4).value, "Noturno")
        self.assertEqual(ws.cell(row=3, column=5).value, "Noturno")    # Nova coluna Restrição
        self.assertEqual(ws.cell(row=3, column=6).value, "15/05/2026") # Data
        self.assertEqual(ws.cell(row=3, column=7).value, "0307968000") # IPTU
        self.assertEqual(ws.cell(row=3, column=8).value, "CENTRAL")     # ÁREA
        self.assertEqual(ws.cell(row=3, column=9).value, "RUA AIMORES, 287 - VILA ANCHIETA")

        # Linha 4 (Restrição ativa, tem dono 'Fiscal_Beta' e não é noturno -> Tipo de Serviço = 'Diurno Livre')
        self.assertEqual(ws.cell(row=4, column=1).value, 2)  # Seq 2
        self.assertIn(ws.cell(row=4, column=2).value, ("", None)) # Status vazio (distribuído)
        self.assertEqual(ws.cell(row=4, column=3).value, "Fiscal_Beta")
        self.assertEqual(ws.cell(row=4, column=4).value, "Diurno Livre")
        self.assertEqual(ws.cell(row=4, column=5).value, "Com Restrição") # Nova coluna Restrição
        self.assertEqual(ws.cell(row=4, column=6).value, "16/05/2026")
        self.assertEqual(ws.cell(row=4, column=7).value, "0554051005")
        self.assertEqual(ws.cell(row=4, column=8).value, "HB")
        self.assertEqual(ws.cell(row=4, column=9).value, "AV MARIO ANDREAZZA, 181 SALA 1 - JARDIM SÃO MARCO")

        # Formatação Condicional presente (15 regras, incluindo Sem Restrição amarelo)
        self.assertEqual(len(ws.conditional_formatting), 1)
        cf_rules = list(ws.conditional_formatting)[0].rules
        self.assertEqual(len(cf_rules), 15)
        self.assertEqual(cf_rules[0].formula[0], '$B3="Em Andamento"')
        self.assertEqual(cf_rules[8].formula[0], '$C3<>""')
        self.assertEqual(cf_rules[9].formula[0], '$D3="Noturno Livre"')
        self.assertEqual(cf_rules[10].formula[0], '$E3="Noturno"')

    def test_status_distribuido_e_pendente_fica_vazio(self) -> None:
        """Garante que não deve ser escrito não distribuído e/ou distribuído, apenas deixe vazio."""
        dados_status = [
            {"protocolo_licenca": "P_PEND", "status": "PENDENTE"},
            {"protocolo_licenca": "P_DIST", "status": "DISTRIBUIDO"},
            {"protocolo_licenca": "P_DIST_ACENTO", "status": "DISTRIBUÍDO"},
            {"protocolo_licenca": "P_NAO_DIST", "status": "NÃO DISTRIBUÍDO"},
            {"protocolo_licenca": "P_NAO_DIST2", "status": "NAO DISTRIBUIDO"},
            {"protocolo_licenca": "P_DEF", "status": "Deferido"},
            {"protocolo_licenca": "P_INDEF", "status": "Indeferido"},
            {"protocolo_licenca": "P_ANDAM", "status": "Em Andamento"},
        ]

        gerar_planilha_trabalho(
            registros_backlog=dados_status,
            caminho_saida=self.output_path,
            caminho_template=self.template_path,
        )

        wb_out = openpyxl.load_workbook(str(self.output_path))
        ws = wb_out["Controle VRE"]

        # Linha 3 a 7: PENDENTE, DISTRIBUIDO, DISTRIBUÍDO, NÃO DISTRIBUÍDO, NAO DISTRIBUIDO devem ser vazios (None)
        for r in range(3, 8):
            val = ws.cell(row=r, column=2).value
            self.assertIn(val, ("", None), f"Linha {r} deveria estar vazia, mas contém: {val!r}")

        # Linhas 8, 9, 10: Status reais de fiscalização devem ser preservados
        self.assertEqual(ws.cell(row=8, column=2).value, "Deferido")
        self.assertEqual(ws.cell(row=8, column=4).value, "Sem Restrição Livre")
        self.assertEqual(ws.cell(row=8, column=5).value, "Sem Restrição")

        # Linha 9 (Indeferido): NÃO deve preencher Tipo de Serviço (Col 4) nem Restrição (Col 5)
        self.assertEqual(ws.cell(row=9, column=2).value, "Indeferido")
        self.assertIn(ws.cell(row=9, column=4).value, ("", None))
        self.assertIn(ws.cell(row=9, column=5).value, ("", None))

        self.assertEqual(ws.cell(row=10, column=2).value, "Em Andamento")
        self.assertEqual(ws.cell(row=10, column=4).value, "Sem Restrição Livre")
        self.assertEqual(ws.cell(row=10, column=5).value, "Sem Restrição")

    def test_sem_restricao_com_cor_amarela(self) -> None:
        """Garante que 'Sem Restrição' possui regra de formatação condicional com amarelo (FFFFC000)."""
        wb = openpyxl.load_workbook(str(self.template_path))
        ws = wb["Controle VRE"]

        encontrou_d = False
        encontrou_e = False
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                for f in rule.formula:
                    if '$D3="Sem Restrição Livre"' in f:
                        encontrou_d = True
                        self.assertEqual(rule.dxf.fill.fgColor.rgb, "FFFFC000")
                    if '$E3="Sem Restrição"' in f:
                        encontrou_e = True
                        self.assertEqual(rule.dxf.fill.fgColor.rgb, "FFFFC000")

        self.assertTrue(encontrou_d, "Regra para Sem Restrição Livre deve estar presente")
        self.assertTrue(encontrou_e, "Regra para Sem Restrição deve estar presente")

    def test_criar_template_padrao_nao_sobrescreve_existente(self) -> None:
        """Verifica que criar_template_padrao não sobrescreve template customizado pelo usuário."""
        wb = openpyxl.load_workbook(str(self.template_path))
        ws = wb["Controle VRE"]
        ws.cell(row=2, column=1, value="Custom Header")
        wb.save(str(self.template_path))

        # Chamada sem forcar não deve sobrescrever
        criar_template_padrao(self.template_path, forcar=False)
        wb_check = openpyxl.load_workbook(str(self.template_path))
        self.assertEqual(wb_check["Controle VRE"].cell(row=2, column=1).value, "Custom Header")

        # Chamada com forcar=True sobrescreve
        criar_template_padrao(self.template_path, forcar=True)
        wb_check2 = openpyxl.load_workbook(str(self.template_path))
        self.assertEqual(wb_check2["Controle VRE"].cell(row=2, column=1).value, "Número ordem ")

    def test_preserva_formatacao_condicional_customizada_do_template(self) -> None:
        """Verifica que regras customizadas inseridas no template são 100% preservadas na geração."""
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import PatternFill

        wb = openpyxl.load_workbook(str(self.template_path))
        ws = wb["Controle VRE"]
        ws.conditional_formatting._cf_rules.clear()

        # Adiciona uma regra customizada única: cor azul para $B3="CUSTOM"
        fill_custom = PatternFill(start_color="FF0000FF", end_color="FF0000FF", fill_type="solid")
        rule_custom = FormulaRule(formula=['$B3="CUSTOM"'], fill=fill_custom)
        ws.conditional_formatting.add("A3:AF5000", rule_custom)
        wb.save(str(self.template_path))

        dados = [{"protocolo_licenca": "P_CUST", "status": "CUSTOM"}]
        gerar_planilha_trabalho(
            registros_backlog=dados,
            caminho_saida=self.output_path,
            caminho_template=self.template_path,
        )

        wb_out = openpyxl.load_workbook(str(self.output_path))
        ws_out = wb_out["Controle VRE"]
        cfs = list(ws_out.conditional_formatting)
        self.assertEqual(len(cfs), 1)
        rules = cfs[0].rules
        self.assertEqual(len(rules), 1)
        self.assertEqual(rules[0].formula[0], '$B3="CUSTOM"')

    def test_preserva_estilos_e_altura_do_template(self) -> None:
        """Verifica que altura e formatação de célula na linha 3 do template são propagadas."""
        from openpyxl.styles import Font

        wb = openpyxl.load_workbook(str(self.template_path))
        ws = wb["Controle VRE"]
        ws.row_dimensions[3].height = 20.0
        ws.cell(row=3, column=1).font = Font(name="Calibri", size=14, bold=True)
        wb.save(str(self.template_path))

        dados = [{"protocolo_licenca": "P_ESTILO"}]
        gerar_planilha_trabalho(
            registros_backlog=dados,
            caminho_saida=self.output_path,
            caminho_template=self.template_path,
        )

        wb_out = openpyxl.load_workbook(str(self.output_path))
        ws_out = wb_out["Controle VRE"]
        self.assertEqual(ws_out.row_dimensions[3].height, 20.0)
        c1 = ws_out.cell(row=3, column=1)
        self.assertEqual(c1.font.name, "Calibri")
        self.assertEqual(c1.font.size, 14)
        self.assertTrue(c1.font.bold)


if __name__ == "__main__":
    unittest.main()
