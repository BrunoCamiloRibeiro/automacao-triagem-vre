"""Testes unitários e de integração para o importador de histórico de planilhas antigas."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.core.database import get_cep_db, get_triagem_db, initialize_schema_if_needed
from src.core.historico_importer import (
    escanear_e_importar_planilhas_antigas,
    importar_planilha_historico,
)


class TestHistoricoImporter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)

        self.triagem_db_path = self.base_path / "triagem_test.db"
        self.cep_db_path = self.base_path / "cep_test.db"

        conn_t = get_triagem_db(self.triagem_db_path)
        initialize_schema_if_needed(conn_t, is_cep=False)
        conn_t.close()

        conn_c = get_cep_db(self.cep_db_path)
        initialize_schema_if_needed(conn_c, is_cep=True)
        conn_c.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_importar_planilha_dinamica(self):
        # Cria uma planilha de teste com colunas arbitrárias
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Controle 2024"

        # Linha 2 com cabeçalhos
        ws.cell(row=2, column=1, value="ID")
        ws.cell(row=2, column=2, value="Fiscal")
        ws.cell(row=2, column=3, value="Data")
        ws.cell(row=2, column=4, value="Área")
        ws.cell(row=2, column=5, value="Cnpj")
        ws.cell(row=2, column=6, value="Bairro")
        ws.cell(row=2, column=7, value="Cep")

        # Dados
        ws.cell(row=3, column=1, value=1)
        ws.cell(row=3, column=2, value="Fiscal Teste A")
        ws.cell(row=3, column=3, value="2024-05-10")
        ws.cell(row=3, column=4, value="Área 2")
        ws.cell(row=3, column=5, value="12.345.678/0001-90")
        ws.cell(row=3, column=6, value="Jd Teste")
        ws.cell(row=3, column=7, value="15000-111")

        ws.cell(row=4, column=1, value=2)
        ws.cell(row=4, column=2, value="Fiscal Teste B")
        ws.cell(row=4, column=3, value="2024-06-15")
        ws.cell(row=4, column=4, value="Área 3")
        ws.cell(row=4, column=5, value="98.765.432/0001-10")
        ws.cell(row=4, column=6, value="Centro")
        ws.cell(row=4, column=7, value="15010-000")

        caminho_xlsx = self.base_path / "planilha_antiga_teste.xlsx"
        wb.save(caminho_xlsx)

        res = importar_planilha_historico(
            caminho_xlsx=caminho_xlsx,
            triagem_db_path=self.triagem_db_path,
            cep_db_path=self.cep_db_path,
            forcar=True,
        )

        self.assertEqual(res["status"], "SUCESSO")
        self.assertEqual(res["fiscais_inseridos"], 2)
        self.assertEqual(res["ceps_inseridos"], 2)
        self.assertEqual(res["protocolos_inseridos"], 2)
        self.assertIn("2024-05", res["meses_identificados"])
        self.assertIn("2024-06", res["meses_identificados"])

        # Verifica persistência no SQLite
        conn_t = get_triagem_db(self.triagem_db_path)
        cur = conn_t.cursor()
        cur.execute("SELECT fiscal FROM historico_fiscais WHERE cnpj_limpo = '12345678000190';")
        row = cur.fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(row["fiscal"], "Fiscal Teste A")

        # Verifica que os protocolos foram adicionados ao vre_backlog e competências registradas
        cur.execute("SELECT COUNT(*) FROM vre_backlog;")
        self.assertEqual(cur.fetchone()[0], 2)

        cur.execute("SELECT COUNT(*) FROM controle_competencias;")
        self.assertEqual(cur.fetchone()[0], 2)
        conn_t.close()

        conn_c = get_cep_db(self.cep_db_path)
        cur_c = conn_c.cursor()
        cur_c.execute("SELECT area FROM cep_mapeamento WHERE cep_limpo = '15000111';")
        row_c = cur_c.fetchone()
        self.assertIsNotNone(row_c)
        self.assertEqual(row_c["area"], "ÁREA 2")
        conn_c.close()

    def test_deduplicacao_linhas_desorganizadas(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Desorganizada"

        # Linha 1: Título solto
        ws.cell(row=1, column=1, value="CONTROLE HISTORICO")

        # Linha 2: Cabeçalho
        ws.cell(row=2, column=1, value="Nº")
        ws.cell(row=2, column=2, value="Fiscal")
        ws.cell(row=2, column=3, value="Data")
        ws.cell(row=2, column=4, value="Área")
        ws.cell(row=2, column=5, value="CNPJ")
        ws.cell(row=2, column=6, value="Bairro")
        ws.cell(row=2, column=7, value="CEP")

        # Linha 3: Dado A (data 2023)
        ws.cell(row=3, column=2, value="Fiscal Silva")
        ws.cell(row=3, column=3, value="10/05/2023")
        ws.cell(row=3, column=4, value="Área 1")
        ws.cell(row=3, column=5, value="11.222.333/0001-99")
        ws.cell(row=3, column=6, value="Centro")
        ws.cell(row=3, column=7, value="15010-000")

        # Linha 4: DUPLICATA exata de A
        ws.cell(row=4, column=2, value="Fiscal Silva")
        ws.cell(row=4, column=3, value="10/05/2023")
        ws.cell(row=4, column=4, value="Área 1")
        ws.cell(row=4, column=5, value="11.222.333/0001-99")
        ws.cell(row=4, column=6, value="Centro")
        ws.cell(row=4, column=7, value="15010-000")

        # Linha 5: Cabeçalho repetido no meio da tabela
        ws.cell(row=5, column=2, value="Fiscal")
        ws.cell(row=5, column=5, value="CNPJ")

        # Linha 6: Linha com termo de pendência / fiscal inválido
        ws.cell(row=6, column=2, value="PENDENTE")
        ws.cell(row=6, column=5, value="55.666.777/0001-88")

        # Linha 7: Dado A com data mais recente (2024) e CNPJ float do Excel (.0)
        ws.cell(row=7, column=2, value="Fiscal Silva")
        ws.cell(row=7, column=3, value="25/11/2024")
        ws.cell(row=7, column=4, value="Área 1")
        ws.cell(row=7, column=5, value=11222333000199.0)

        # Linha 8: Empresa B
        ws.cell(row=8, column=2, value="Fiscal Souza")
        ws.cell(row=8, column=3, value="12/01/2024")
        ws.cell(row=8, column=4, value="Área 4")
        ws.cell(row=8, column=5, value="33.444.555/0001-22")
        ws.cell(row=8, column=6, value="Redentora")
        ws.cell(row=8, column=7, value="15015-000")

        caminho_xlsx = self.base_path / "planilha_baguncada.xlsx"
        wb.save(caminho_xlsx)

        res = importar_planilha_historico(
            caminho_xlsx=caminho_xlsx,
            triagem_db_path=self.triagem_db_path,
            cep_db_path=self.cep_db_path,
            forcar=True,
        )

        self.assertEqual(res["status"], "SUCESSO")
        self.assertEqual(res["fiscais_inseridos"], 2)
        self.assertEqual(res["duplicatas_descartadas"], 2)
        self.assertEqual(res["ceps_inseridos"], 2)

        # Verifica no SQLite se a data mais recente (2024-11-25) prevaleceu
        conn_t = get_triagem_db(self.triagem_db_path)
        cur = conn_t.cursor()
        cur.execute("SELECT fiscal, data_atendimento FROM historico_fiscais WHERE cnpj_limpo = '11222333000199';")
        rows = cur.fetchall()
        self.assertEqual(len(rows), 1)  # estritamente 1 registro sem duplicata
        self.assertEqual(rows[0]["fiscal"], "Fiscal Silva")
        self.assertEqual(rows[0]["data_atendimento"], "2024-11-25")
        conn_t.close()

    def test_duas_colunas_data_prioriza_data_vre(self):
        """Verifica que, quando há duas colunas de Data (colocada e VRE), deve-se usar a data do VRE."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DuasDatas"

        # Linha 2: Cabeçalho com 2 colunas de Data
        # Col 2: Fiscal, Col 3: Data (colocada), Col 5: CNPJ, Col 6: ProtocoloLicenca, Col 7: Data (VRE)
        ws.cell(row=2, column=1, value="Nº")
        ws.cell(row=2, column=2, value="Fiscal")
        ws.cell(row=2, column=3, value="Data")               # Data colocada
        ws.cell(row=2, column=4, value="Área")
        ws.cell(row=2, column=5, value="CNPJ")
        ws.cell(row=2, column=6, value="ProtocoloLicenca")
        ws.cell(row=2, column=7, value="Data")               # Data do VRE

        # Linha 3: Protocolo com data colocada em junho/2026, mas data do VRE em janeiro/2026
        ws.cell(row=3, column=1, value=1)
        ws.cell(row=3, column=2, value="Fiscal_Alfa")
        ws.cell(row=3, column=3, value="17/06/2026")         # Data colocada na distribuição
        ws.cell(row=3, column=4, value="HB")
        ws.cell(row=3, column=5, value="33.424.347/0001-10")
        ws.cell(row=3, column=6, value="LIC_VRE_001")
        ws.cell(row=3, column=7, value="02/01/2026")         # Data oficial do VRE

        caminho_xlsx = self.base_path / "planilha_duas_datas.xlsx"
        wb.save(caminho_xlsx)

        res = importar_planilha_historico(
            caminho_xlsx=caminho_xlsx,
            triagem_db_path=self.triagem_db_path,
            cep_db_path=self.cep_db_path,
            forcar=True,
        )

        self.assertEqual(res["status"], "SUCESSO")
        self.assertIn("2026-01", res["meses_identificados"])
        self.assertNotIn("2026-06", res["meses_identificados"])

        conn_t = get_triagem_db(self.triagem_db_path)
        cur = conn_t.cursor()
        cur.execute("SELECT mes_competencia, data_importacao FROM vre_backlog WHERE protocolo_licenca = 'LIC_VRE_001';")
        r = cur.fetchone()
        self.assertIsNotNone(r)
        self.assertEqual(r["mes_competencia"], "2026-01", "Deve usar a competência do VRE (janeiro), não a colocada (junho).")
        self.assertEqual(r["data_importacao"], "2026-01-02", "Deve usar a data do VRE, não a colocada.")
        conn_t.close()


if __name__ == "__main__":
    unittest.main()
