"""Teste de Integração Completo de Ponta a Ponta (End-to-End).

Simula todo o ecossistema com dados sintéticos abrangendo todos os casos de borda:
1. Resolução de Fiscal por CNPJ histórico.
2. Atribuição de Flags e Tipo de Serviço para CNAEs (Noturno, Restrição, Ambos).
3. Resolução de Área por Bairro normalizado.
4. Resolução de Área via Fallback para CEP (banco cep_areas.db).
5. Registro de CEP Pendente quando nem bairro nem CEP forem conhecidos.
6. Proteção de estado prévio no backlog (INSERT OR IGNORE).
7. Geração de Planilha com Plain Values, formatação condicional, painéis e filtros.
8. Transação atômica de arquivamento de competências.
"""

from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.admin.archive_service import arquivar_competencia, listar_competencias_status
from src.core.cep_service import consultar_area_por_cep, listar_ceps_pendentes, registrar_area_cep
from src.core.database import connect_db, initialize_schema_if_needed
from src.core.excel_writer import criar_template_padrao
from src.runner.main_seco import processar_fluxo_seco


class TestIntegracaoCompleta(unittest.TestCase):
    """Teste de integração ponta a ponta em ambiente controlado e isolado."""

    def setUp(self) -> None:
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp_dir.name)

        # Caminhos dos bancos e arquivos temporários
        self.db_triagem = self.root / "triagem_vre.db"
        self.db_cep = self.root / "cep_areas.db"
        self.template_xlsx = self.root / "template_triagem.xlsx"
        self.saida_xlsx = self.root / "Planilha_Controle_VRE_Ativa.xlsx"
        self.csv_entrada = self.root / "remessa_teste.csv"

        # Inicializa schemas DDL
        self.conn_t = connect_db(self.db_triagem)
        initialize_schema_if_needed(self.conn_t, is_cep=False)

        self.conn_c = connect_db(self.db_cep)
        initialize_schema_if_needed(self.conn_c, is_cep=True)

        # Gera o template mestre
        criar_template_padrao(self.template_xlsx)

        # 1. Popula Histórico de Fiscais prévio
        cur_t = self.conn_t.cursor()
        cur_t.execute(
            """
            INSERT INTO historico_fiscais (cnpj_limpo, fiscal, data_atendimento, origem_registro)
            VALUES ('11111111000111', 'Fiscal_Alfa', '2025-10-15', 'Legado 2025');
            """
        )

        # 2. Popula Regras de CNAEs
        # CNAE A: 5611-2/01 -> Noturno = 1
        # CNAE B: 4731-8/00 -> Restrição = 1
        # CNAE C: 4723-7/00 -> Noturno = 1
        # CNAE D: 0111-3/01 -> Livre (0, 0)
        cur_t.executemany(
            """
            INSERT INTO regras_cnae (cnae_7digitos, cnae_formatado, denominacao, restricao, noturno)
            VALUES (?, ?, ?, ?, ?);
            """,
            [
                ("5611201", "5611-2/01", "Restaurantes e similares", 0, 1),
                ("4731800", "4731-8/00", "Comércio de Combustíveis", 1, 0),
                ("4723700", "4723-7/00", "Comércio de Bebidas", 0, 1),
                ("0111301", "0111-3/01", "Cultivo de Arroz", 0, 0),
            ],
        )

        # 3. Popula Bairros x Áreas
        cur_t.executemany(
            """
            INSERT INTO bairros_areas (bairro_normalizado, area)
            VALUES (?, ?);
            """,
            [
                ("VILA ANCHIETA", "REPRESA"),
                ("JARDIM MARACANA", "BOSQUE"),
            ],
        )

        # 4. Popula CEPs Mapeados em cep_areas.db
        registrar_area_cep(
            cep="15080000",
            area="HB",
            bairro="Bairro Inconclusivo",
            origem="MANUAL",
            conn=self.conn_c,
        )

        self.conn_t.close()
        self.conn_c.close()

    def tearDown(self) -> None:
        self.tmp_dir.cleanup()

    def test_pipeline_ponta_a_ponta(self) -> None:
        # Monta um CSV sintético com múltiplos casos
        cabecalhos = [
            "ProtocoloSil",
            "ProtocoloLicenca",
            "Data",
            "GrauRiscoGeral",
            "SituacaoPostura",
            "InscricaoMunicipal",
            "Cnpj",
            "NomeEmpresarial",
            "Telefone",
            "E-mail",
            "TipoEndereco",
            "Logradouro",
            "Numero",
            "Complemento",
            "Bairro",
            "Cep",
            "FinalidadeEndereco",
            "CnaesLicenciadas",
            "CnaesNaoLicenciadas",
            "Iptu",
        ]

        linhas_teste = [
            # Caso 1: CNPJ com fiscal histórico prévio + CNAE Noturno + Bairro conhecido
            [
                "SIL_001",
                "LIC_001",
                "10/06/2026",
                "BAIXO",
                "DEFERIDO",
                "1001",
                "11.111.111/0001-11",
                "RESTAURANTE MODELO LTDA",
                "17-99991111",
                "contato@restaurante.com",
                "RUA",
                "AIMORES",
                "100",
                ":",
                "Vila Anchieta",
                "15050090",
                "EXERCICIO DA ATIVIDADE",
                "561120101",  # 5611201 -> Noturno
                "",
                "010203000",
            ],
            # Caso 2: CNPJ novo + CNAE Restrição + Bairro desconhecido, mas CEP mapeado (Fallback)
            [
                "SIL_002",
                "LIC_002",
                "11/06/2026",
                "ALTO",
                "DEFERIDO",
                "1002",
                "22.222.222/0001-22",
                "POSTO COMBUSTIVEL LTDA",
                "17-99992222",
                "posto@combustivel.com",
                "AV",
                "BADY BASSITT",
                "500",
                "BLOCO A",
                "Bairro Desconhecido Sem Mapeamento",
                "15080-000",  # CEP mapeado para Área 'HB'
                "EXERCICIO DA ATIVIDADE",
                "473180001",  # 4731800 -> Restrição
                "",
                "020304000",
            ],
            # Caso 3: Ambos os flags (Noturno + Restrição via curto-circuito) + Nem bairro nem CEP mapeados
            [
                "SIL_003",
                "LIC_003",
                "12/06/2026",
                "ALTO",
                "DEFERIDO",
                "1003",
                "33.333.333/0001-33",
                "BEBIDAS E COMBUSTIVEIS LTDA",
                "17-99993333",
                "contato@bebidas.com",
                "RUA",
                "CENTRAL",
                "99",
                "",
                "Bairro Fantasma",
                "15099-999",  # CEP desconhecido
                "EXERCICIO DA ATIVIDADE",
                "473180001, 472370001",  # 4731800 (Restrição=1) + 4723700 (Noturno=1)
                "",
                "030405000",
            ],
        ]

        # Escreve o CSV em Latin-1 com separador ';'
        with open(self.csv_entrada, "w", encoding="latin-1", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(cabecalhos)
            writer.writerows(linhas_teste)

        # Executa o runner seco apontando para nossos bancos e template isolados
        import src.core.config as cfg

        # Ajusta caminhos do config temporariamente para o teste
        orig_triagem = cfg.TRIAGEM_DB_PATH
        orig_cep = cfg.CEP_DB_PATH
        orig_template = cfg.TEMPLATE_XLSX_PATH
        orig_saida = cfg.OUTPUT_XLSX_PATH

        try:
            cfg.TRIAGEM_DB_PATH = self.db_triagem
            cfg.CEP_DB_PATH = self.db_cep
            cfg.TEMPLATE_XLSX_PATH = self.template_xlsx
            cfg.OUTPUT_XLSX_PATH = self.saida_xlsx

            total_ativos = processar_fluxo_seco(
                caminho_csv=self.csv_entrada,
                saida_xlsx=self.saida_xlsx,
                silencioso=True,
                triagem_db_path=self.db_triagem,
                cep_db_path=self.db_cep,
                template_path=self.template_xlsx,
            )
            self.assertEqual(total_ativos, 3)

            # =================================================================
            # VERIFICAÇÃO 1: DADOS NO BANCO triagem_vre.db
            # =================================================================
            conn_t = connect_db(self.db_triagem)
            cur_t = conn_t.cursor()

            # Registro 1
            cur_t.execute("SELECT * FROM vre_backlog WHERE protocolo_licenca = 'LIC_001';")
            r1 = cur_t.fetchone()
            self.assertEqual(r1["fiscal_atribuido"], "Fiscal_Alfa", "Deve herdar o fiscal histórico atribuído.")
            self.assertEqual(r1["flag_noturno"], 1, "Deve marcar Noturno=1.")
            self.assertEqual(r1["flag_restricao"], 0, "Deve manter Restrição=0.")
            self.assertEqual(r1["tipo_servico"], "Noturno")
            self.assertEqual(r1["area"], "REPRESA", "Deve resolver Área pelo Bairro normalizado.")

            # Registro 2
            cur_t.execute("SELECT * FROM vre_backlog WHERE protocolo_licenca = 'LIC_002';")
            r2 = cur_t.fetchone()
            self.assertEqual(r2["fiscal_atribuido"], "", "Fiscal deve ser vazio.")
            self.assertEqual(r2["flag_noturno"], 0)
            self.assertEqual(r2["flag_restricao"], 1)
            self.assertEqual(r2["tipo_servico"], "Com Restrição Livre")
            self.assertEqual(r2["area"], "HB", "Deve resolver Área via Fallback para cep_areas.db.")

            # Registro 3
            cur_t.execute("SELECT * FROM vre_backlog WHERE protocolo_licenca = 'LIC_003';")
            r3 = cur_t.fetchone()
            self.assertEqual(r3["flag_noturno"], 1)
            self.assertEqual(r3["flag_restricao"], 1)
            self.assertEqual(r3["tipo_servico"], "Noturno Livre")
            self.assertEqual(r3["area"], "", "Área deve ficar vazia quando inconclusiva.")

            conn_t.close()

            # =================================================================
            # VERIFICAÇÃO 2: CEP PENDENTE EM cep_areas.db
            # =================================================================
            conn_c = connect_db(self.db_cep)
            pendentes = listar_ceps_pendentes(conn=conn_c)
            self.assertTrue(
                any(p["cep_limpo"] == "15099999" for p in pendentes),
                "O CEP desconhecido 15099999 deve ser registrado em cep_pendentes.",
            )
            conn_c.close()

            # =================================================================
            # VERIFICAÇÃO 3: PROTEÇÃO CONTRA SOBREESCRITA (INSERT OR IGNORE)
            # =================================================================
            # Simula o operador editando manualmente o status e fiscal do LIC_001
            conn_t = connect_db(self.db_triagem)
            conn_t.execute(
                """
                UPDATE vre_backlog
                SET fiscal_atribuido = 'Fiscal_Designado_Manual',
                    status = 'DISTRIBUIDO'
                WHERE protocolo_licenca = 'LIC_001';
                """
            )
            conn_t.close()

            # Roda novamente a importação do mesmo CSV
            processar_fluxo_seco(
                caminho_csv=self.csv_entrada,
                saida_xlsx=self.saida_xlsx,
                silencioso=True,
                triagem_db_path=self.db_triagem,
                cep_db_path=self.db_cep,
                template_path=self.template_xlsx,
            )

            conn_t = connect_db(self.db_triagem)
            cur_t = conn_t.cursor()
            cur_t.execute("SELECT fiscal_atribuido, status FROM vre_backlog WHERE protocolo_licenca = 'LIC_001';")
            r1_pos = cur_t.fetchone()
            self.assertEqual(r1_pos["fiscal_atribuido"], "Fiscal_Designado_Manual", "INSERT OR IGNORE deve proteger o fiscal alterado.")
            self.assertEqual(r1_pos["status"], "DISTRIBUIDO", "INSERT OR IGNORE deve proteger o status alterado.")
            conn_t.close()

            # =================================================================
            # VERIFICAÇÃO 4: PLANILHA EXCEL GERADA
            # =================================================================
            self.assertTrue(self.saida_xlsx.exists())
            wb = openpyxl.load_workbook(str(self.saida_xlsx))
            ws = wb["Controle VRE"]

            self.assertEqual(ws.freeze_panes, "A3")
            self.assertEqual(ws.max_row, 5)  # Linha 1 cabeçalho controle, Linha 2 cabeçalhos, Linhas 3, 4, 5 dados
            self.assertEqual(ws.max_column, 32)

            # Linha 3 (LIC_001 - Noturno)
            self.assertEqual(ws.cell(row=3, column=4).value, "Noturno")
            self.assertEqual(ws.cell(row=3, column=5).value, "Noturno")
            self.assertEqual(ws.cell(row=3, column=8).value, "REPRESA")

            # Linha 4 (LIC_002 - Restrição)
            self.assertEqual(ws.cell(row=4, column=4).value, "Com Restrição Livre")
            self.assertEqual(ws.cell(row=4, column=5).value, "Com Restrição")
            self.assertEqual(ws.cell(row=4, column=8).value, "HB")

            # Linha 5 (LIC_003 - Ambos CNAEs, área desconhecida)
            self.assertEqual(ws.cell(row=5, column=4).value, "Noturno Livre")
            self.assertEqual(ws.cell(row=5, column=5).value, "Noturno")
            self.assertIn(ws.cell(row=5, column=8).value, ("", None))    # Área vazia

            # =================================================================
            # VERIFICAÇÃO 5: TRANSAÇÃO DE ARQUIVAMENTO DE COMPETÊNCIAS
            # =================================================================
            # Marca todos os protocolos da competência 2026-06 como FINALIZADO
            conn_t = connect_db(self.db_triagem)
            conn_t.execute(
                """
                UPDATE vre_backlog
                SET status = 'FINALIZADO',
                    fiscal_atribuido = CASE
                        WHEN fiscal_atribuido IS NULL OR TRIM(fiscal_atribuido) = '' THEN 'Fiscal SMAURB'
                        ELSE fiscal_atribuido
                    END,
                    data_finalizacao = '2026-06-30'
                WHERE mes_competencia = '2026-06';
                """
            )

            # Valida status do lote
            lotes = listar_competencias_status(conn=conn_t)
            self.assertEqual(len(lotes), 1)
            self.assertTrue(lotes[0]["pode_arquivar"])

            # Executa o arquivamento definitivo
            qtd_arq, msg_arq = arquivar_competencia("2026-06", forcar=False, conn=conn_t)
            self.assertEqual(qtd_arq, 3)

            # Verifica que os 3 registros permanecem preservados com segurança no banco SQLite
            cur_t = conn_t.cursor()
            cur_t.execute("SELECT COUNT(*) FROM vre_backlog;")
            self.assertEqual(cur_t.fetchone()[0], 3)

            # Verifica se foram parar no historico_fiscais
            cur_t.execute("SELECT COUNT(*) FROM historico_fiscais WHERE origem_registro = 'Mes_2026-06';")
            self.assertEqual(cur_t.fetchone()[0], 3)

            conn_t.close()

            # Re-gera a planilha e verifica que ela está vazia de linhas de dados (apenas cabeçalhos)
            total_apos_arq = processar_fluxo_seco(
                caminho_csv="",
                saida_xlsx=self.saida_xlsx,
                silencioso=True,
                triagem_db_path=self.db_triagem,
                cep_db_path=self.db_cep,
                template_path=self.template_xlsx,
            )
            self.assertEqual(total_apos_arq, 0)

            wb_vazio = openpyxl.load_workbook(str(self.saida_xlsx))
            ws_vazio = wb_vazio["Controle VRE"]
            self.assertEqual(ws_vazio.max_row, 2, "A planilha ativa deve conter apenas os cabeçalhos após arquivamento.")

        finally:
            # Restaura caminhos de configuração originais
            cfg.TRIAGEM_DB_PATH = orig_triagem
            cfg.CEP_DB_PATH = orig_cep
            cfg.TEMPLATE_XLSX_PATH = orig_template
            cfg.OUTPUT_XLSX_PATH = orig_saida

    def test_preservacao_edicoes_manuais_na_planilha_ativa(self) -> None:
        """
        Garante que:
        1. Quando o usuário sobe o CSV do Dia 1, edita a Planilha_Controle_VRE_Ativa.xlsx
           (ex: Status = Deferido, Em Andamento, Fiscal, ProtocoloSigm, IPTU, Área)
           e depois sobe o CSV do Dia 2, TODAS as linhas e modificações manuais do Dia 1
           permanecem intactas na planilha gerada no Dia 2.
        2. Registros INDEFERIDOS (novos ou passados) ficam sem Tipo de Serviço e sem Restrição.
        3. A Coluna F (Data) permanece vazia por padrão e preserva apenas datas preenchidas a mão.
        """
        import src.core.config as cfg

        orig_triagem = cfg.TRIAGEM_DB_PATH
        orig_cep = cfg.CEP_DB_PATH
        orig_template = cfg.TEMPLATE_XLSX_PATH
        orig_saida = cfg.OUTPUT_XLSX_PATH

        csv_dia1 = self.root / "dia1.csv"
        csv_dia2 = self.root / "dia2.csv"

        cabecalho = [
            "ProtocoloSil", "ProtocoloLicenca", "Data", "GrauRiscoGeral",
            "SituacaoPostura", "InscricaoMunicipal", "Cnpj", "NomeEmpresarial",
            "Telefone", "E-mail", "TipoEndereco", "Logradouro", "Numero",
            "Complemento", "Bairro", "Cep", "FinalidadeEndereco",
            "CnaesLicenciadas", "CnaesNaoLicenciadas", "Iptu",
        ]

        with open(csv_dia1, mode="w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(cabecalho)
            writer.writerows([
                [
                    "SIL_D1_1", "LIC_D1_1", "01/06/2026", "MEDIO", "DEFERIDO", "2001",
                    "44.444.444/0001-44", "PADARIA DIA UM LTDA", "17-99994444", "d1@teste.com",
                    "RUA", "UM", "10", "", "Vila Anchieta", "15050090",
                    "EXERCICIO DA ATIVIDADE", "561120101", "", "010101000",
                ],
                [
                    "SIL_D1_2", "LIC_D1_2", "01/06/2026", "ALTO", "PENDENTE", "2002",
                    "55.555.555/0001-55", "MERCADO DIA UM S/A", "17-99995555", "d1b@teste.com",
                    "AV", "DOIS", "20", "", "Bairro Inexistente", "15081490",
                    "EXERCICIO DA ATIVIDADE", "461410004", "", "",
                ],
            ])

        with open(csv_dia2, mode="w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(cabecalho)
            writer.writerows([
                [
                    "SIL_D2_1", "LIC_D2_1", "02/06/2026", "BAIXO", "DEFERIDO", "2003",
                    "66.666.666/0001-66", "FARMACIA DIA DOIS LTDA", "17-99996666", "d2@teste.com",
                    "RUA", "TRES", "30", "", "Jardim Maracana", "15050100",
                    "EXERCICIO DA ATIVIDADE", "011130101", "", "020202000",
                ],
                [
                    "SIL_D2_INDEF", "LIC_D2_INDEF", "02/06/2026", "ALTO", "INDEFERIDO", "2004",
                    "77.777.777/0001-77", "EMPRESA INDEFERIDA LTDA", "17-99997777", "indef@teste.com",
                    "RUA", "QUATRO", "40", "", "Vila Anchieta", "15050090",
                    "EXERCICIO DA ATIVIDADE", "561120101", "", "030303000",
                ],
            ])

        try:
            cfg.TRIAGEM_DB_PATH = self.db_triagem
            cfg.CEP_DB_PATH = self.db_cep
            cfg.TEMPLATE_XLSX_PATH = self.template_xlsx
            cfg.OUTPUT_XLSX_PATH = self.saida_xlsx

            # Simula um registro passado no banco que estava INDEFERIDO mas tinha Tipo de Serviço e Restrição preenchidos antes do patch
            conn_t = connect_db(self.db_triagem)
            conn_t.execute(
                """
                INSERT INTO vre_backlog (
                    protocolo_licenca, protocolo_sil, cnpj_limpo, cnpj_formatado,
                    nome_empresarial, mes_competencia, data_importacao,
                    bairro_original, area, cnaes_licenciadas, fiscal_atribuido,
                    flag_restricao, flag_noturno, tipo_servico, status, dados_extras_json
                )
                VALUES (
                    'LIC_PASSADO_INDEF', 'SIL_PASS_1', '88888888000188', '88.888.888/0001-88',
                    'EMPRESA PASSADA INDEFERIDA', '2026-06', '2026-06-01',
                    'Vila Anchieta', 'REPRESA', '561120101', '',
                    1, 1, 'Noturno Livre', 'PENDENTE',
                    '{"SituacaoPostura": "INDEFERIDO", "RestricaoCustom": "Noturno", "TipoServicoCustom": "Noturno Livre", "DataCol6": "01/01/2025"}'
                );
                """
            )
            conn_t.close()

            # 1. Sobe o Dia 1 (2026-06-01)
            total_d1 = processar_fluxo_seco(
                caminho_csv=csv_dia1,
                saida_xlsx=self.saida_xlsx,
                silencioso=True,
                triagem_db_path=self.db_triagem,
                cep_db_path=self.db_cep,
                template_path=self.template_xlsx,
            )
            self.assertEqual(total_d1, 3)

            # 2. Operador abre o Excel do Dia 1 e faz modificações manuais
            wb_d1 = openpyxl.load_workbook(str(self.saida_xlsx))
            ws_d1 = wb_d1["Controle VRE"]

            # Mapeia linhas no Dia 1 por ProtocoloLicenca (Coluna 13)
            mapa_d1 = {ws_d1.cell(row=r, column=13).value: r for r in range(3, ws_d1.max_row + 1)}

            # Verifica que o registro passado indeferido já teve Tipo de Serviço, Restrição e DataCol6 (Col 6) limpos no Dia 1
            r_pass_d1 = mapa_d1["LIC_PASSADO_INDEF"]
            self.assertIn(ws_d1.cell(row=r_pass_d1, column=4).value, ("", None))
            self.assertIn(ws_d1.cell(row=r_pass_d1, column=5).value, ("", None))
            self.assertIn(ws_d1.cell(row=r_pass_d1, column=6).value, ("", None))

            # Verifica que todas as linhas no Dia 1 iniciam com a Coluna F (6) vazia
            for r_chk in mapa_d1.values():
                self.assertIn(ws_d1.cell(row=r_chk, column=6).value, ("", None))

            # Edita LIC_D1_1: muda Status para Deferido, Fiscal, Tipo de Serviço para 'Noturno entregue', Data (Col F) a mão, ProtocoloSigm
            r1 = mapa_d1["LIC_D1_1"]
            ws_d1.cell(row=r1, column=2, value="Deferido")
            ws_d1.cell(row=r1, column=3, value="Fiscal_Gama")
            ws_d1.cell(row=r1, column=4, value="Noturno entregue")
            ws_d1.cell(row=r1, column=6, value="18/06/2026")
            ws_d1.cell(row=r1, column=10, value="Sim")
            ws_d1.cell(row=r1, column=11, value="2026/99999")

            # Edita LIC_D1_2: muda Status para Em Andamento, Fiscal, preenche IPTU e ÁREA manualmente
            r2 = mapa_d1["LIC_D1_2"]
            ws_d1.cell(row=r2, column=2, value="Em Andamento")
            ws_d1.cell(row=r2, column=3, value="Fiscal_Delta")
            ws_d1.cell(row=r2, column=7, value="099988800")
            ws_d1.cell(row=r2, column=8, value="CENTRO")
            ws_d1.cell(row=r2, column=11, value="OS-1234")

            wb_d1.save(str(self.saida_xlsx))
            wb_d1.close()

            # 3. Sobe o Dia 2 (2026-06-02)
            total_d2 = processar_fluxo_seco(
                caminho_csv=csv_dia2,
                saida_xlsx=self.saida_xlsx,
                silencioso=True,
                triagem_db_path=self.db_triagem,
                cep_db_path=self.db_cep,
                template_path=self.template_xlsx,
            )
            self.assertEqual(total_d2, 5, "Todos os 5 protocolos devem permanecer na planilha ativa.")

            # 4. Verifica se as modificações do Dia 1 permaneceram intactas na planilha gerada no Dia 2
            wb_d2 = openpyxl.load_workbook(str(self.saida_xlsx))
            ws_d2 = wb_d2["Controle VRE"]
            self.assertEqual(ws_d2.max_row, 7)  # 2 linhas de cabeçalho + 5 registros

            # Mapeia linhas por ProtocoloLicenca (Coluna 13)
            linhas_por_lic = {}
            for r in range(3, ws_d2.max_row + 1):
                lic = ws_d2.cell(row=r, column=13).value
                linhas_por_lic[lic] = r

            self.assertIn("LIC_D1_1", linhas_por_lic)
            self.assertIn("LIC_D1_2", linhas_por_lic)
            self.assertIn("LIC_D2_1", linhas_por_lic)
            self.assertIn("LIC_D2_INDEF", linhas_por_lic)
            self.assertIn("LIC_PASSADO_INDEF", linhas_por_lic)

            r_d1_1 = linhas_por_lic["LIC_D1_1"]
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=2).value, "Deferido")
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=3).value, "Fiscal_Gama")
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=4).value, "Noturno entregue")
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=6).value, "18/06/2026")
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=10).value, "Sim")
            self.assertEqual(ws_d2.cell(row=r_d1_1, column=11).value, "2026/99999")

            r_d1_2 = linhas_por_lic["LIC_D1_2"]
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=2).value, "Em Andamento")
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=3).value, "Fiscal_Delta")
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=4).value, "Diurno Livre")
            self.assertIn(ws_d2.cell(row=r_d1_2, column=6).value, ("", None))
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=7).value, "099988800")
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=8).value, "CENTRO")
            self.assertEqual(ws_d2.cell(row=r_d1_2, column=11).value, "OS-1234")

            # Verifica que novos registros do Dia 2 vêm com Coluna F (6) vazia
            self.assertIn(ws_d2.cell(row=linhas_por_lic["LIC_D2_1"], column=6).value, ("", None))

            # Verifica que tanto o INDEFERIDO novo quanto o INDEFERIDO passado estão sem Tipo de Serviço e sem Restrição
            for lic_indef in ("LIC_D2_INDEF", "LIC_PASSADO_INDEF"):
                r_ind = linhas_por_lic[lic_indef]
                self.assertIn(ws_d2.cell(row=r_ind, column=4).value, ("", None))
                self.assertIn(ws_d2.cell(row=r_ind, column=5).value, ("", None))
                self.assertIn(ws_d2.cell(row=r_ind, column=6).value, ("", None))

            wb_d2.close()

        finally:
            cfg.TRIAGEM_DB_PATH = orig_triagem
            cfg.CEP_DB_PATH = orig_cep
            cfg.TEMPLATE_XLSX_PATH = orig_template
            cfg.OUTPUT_XLSX_PATH = orig_saida


if __name__ == "__main__":
    unittest.main()
